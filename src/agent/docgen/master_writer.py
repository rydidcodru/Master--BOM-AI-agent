"""통합 개발부품 Master(Compact v1.1 형식) 서식보존 출력.

정답지 `통합 개발부품Master Compact v1.1.xlsx`("Master" 시트)를 **템플릿**으로 열어
다중헤더(1~9행)·병합·스타일을 그대로 두고 데이터행(10행~)만 우리 산출로 교체한다.

정답지 컬럼 규약(2026-06-13 역공학):
  B(2) No.  C(3) BOM Level  D(4) Part Type  E(5) Base P/No  F(6) New P/No
  G(7) 부품명  H(8) Q'ty Base  I(9) Q'ty New  J(10) 변경점  K(11) 변경사유
  L(12) 양산처  M(13) 분류(Common/New/Change/Delete)
  * Common  : F='←', I='←', J='-', K='-', M='Common'
  * 변경/신규: F=new pno, M=New|Change (정답지 New=대다수 base→new 번호교체)
  * Delete  : F='X'(완전삭제) 또는 교체품번, M='Delete'

분류(M)는 **진단용**(정답지의 New/Change 구분은 사람 판단이라 결정론 재현 불가 →
절대원칙 #4: 분류는 HITL 확정). 게이팅 신호는 base_pno(E)+new_pno(F) 쌍이다.
LLM/DB/네트워크 0회.
"""
from __future__ import annotations

import io
from dataclasses import dataclass

from src.agent.basebom.diff import DiffRow

__all__ = ["MasterRow", "rows_from_diff", "write_master_xlsx", "DATA_START_ROW", "COLS"]

# 정답지 "Master" 시트 데이터 시작행 / 컬럼(1-indexed).
DATA_START_ROW = 10
COLS = {
    "no": 2, "bom_level": 3, "part_type": 4, "base_pno": 5, "new_pno": 6,
    "part_name": 7, "qty_base": 8, "qty_new": 9, "changing_point": 10,
    "changing_reason": 11, "supplier": 12, "classification": 13,
}
SHEET_NAME = "Master"

# 변경점(J) 거친 3분류 — 정답지의 세분류는 사유 의존이라 진단용 coarse 매핑.
_COARSE_POINT = {"Change": "부품 변경", "New": "부품 추가", "Delete": "부품 삭제", "Common": "-"}


@dataclass
class MasterRow:
    """통합 master 한 데이터행 (정답지 형식)."""
    bom_level: str
    part_type: str
    base_pno: str
    new_pno: str        # '←'=미변경, 'X'=삭제, 그 외=신규/교체 품번
    part_name: str
    qty_base: str
    qty_new: str
    changing_point: str
    changing_reason: str
    supplier: str
    classification: str


def rows_from_diff(
    diff_rows: list[DiffRow],
    *,
    reasons: dict[str, str] | None = None,
    suppliers: dict[str, str] | None = None,
) -> list[MasterRow]:
    """DiffRow 리스트 → MasterRow 리스트 (정답지 셀 규약 적용).

    Args:
        diff_rows: ``diff_boms`` 결과(base 트리 순서 + New append).
        reasons: base_pno(정규화 전 원본) → 변경사유(K). PPT/후보에서 결합한 best-effort.
        suppliers: base_pno → 양산처(L). base BOM Maker 등에서.
    """
    reasons = reasons or {}
    suppliers = suppliers or {}
    out: list[MasterRow] = []
    for d in diff_rows:
        cls = d.classification
        key = d.base_pno or d.new_pno
        if cls == "Common":
            new_pno_cell, qty_new, jp, kr = "←", "←", "-", "-"
        elif cls == "Delete":
            new_pno_cell = d.new_pno or "X"  # 교체품번 있으면 표기, 없으면 완전삭제 'X'
            qty_new, jp = "X", _COARSE_POINT["Delete"]
            kr = reasons.get(key, "-") or "-"
        else:  # Change / New — base→new 번호교체 또는 신규
            new_pno_cell = d.new_pno or "<발번대기>"  # 신규인데 번호 없으면 placeholder(절대원칙)
            qty_new = d.qty or ""
            jp = _COARSE_POINT.get(cls, "부품 변경")
            kr = reasons.get(key, "-") or "-"
        out.append(MasterRow(
            bom_level=d.lvl, part_type=d.part_type, base_pno=d.base_pno,
            new_pno=new_pno_cell, part_name=d.part_name, qty_base=d.qty,
            qty_new=qty_new, changing_point=jp, changing_reason=kr,
            supplier=suppliers.get(key, "-") or "-", classification=cls,
        ))
    return out


def write_master_xlsx(rows: list[MasterRow], template_path: str) -> bytes:
    """MasterRow 리스트 → 정답지 형식 xlsx bytes (헤더 보존, 데이터행 교체).

    템플릿(정답지)을 열어 1~9행 헤더/병합/스타일은 그대로, 10행 이하 데이터셀만
    지우고 우리 산출로 채운다. 기존 정답 데이터는 비교 오염 방지를 위해 선삭제.
    """
    import openpyxl

    wb = openpyxl.load_workbook(template_path, data_only=False)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]

    # 데이터행(10행~) 기존 값 삭제 — 스코어 대상 컬럼만(병합/스타일 보존, 메타컬럼 미건드림).
    scored_cols = list(COLS.values())
    for ri in range(DATA_START_ROW, ws.max_row + 1):
        for ci in scored_cols:
            ws.cell(row=ri, column=ci).value = None

    # 우리 산출 기입.
    for i, r in enumerate(rows):
        ri = DATA_START_ROW + i
        ws.cell(row=ri, column=COLS["no"], value=i + 1)
        ws.cell(row=ri, column=COLS["bom_level"], value=r.bom_level)
        ws.cell(row=ri, column=COLS["part_type"], value=r.part_type)
        ws.cell(row=ri, column=COLS["base_pno"], value=r.base_pno)
        ws.cell(row=ri, column=COLS["new_pno"], value=r.new_pno)
        ws.cell(row=ri, column=COLS["part_name"], value=r.part_name)
        ws.cell(row=ri, column=COLS["qty_base"], value=r.qty_base)
        ws.cell(row=ri, column=COLS["qty_new"], value=r.qty_new)
        ws.cell(row=ri, column=COLS["changing_point"], value=r.changing_point)
        ws.cell(row=ri, column=COLS["changing_reason"], value=r.changing_reason)
        ws.cell(row=ri, column=COLS["supplier"], value=r.supplier)
        ws.cell(row=ri, column=COLS["classification"], value=r.classification)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
