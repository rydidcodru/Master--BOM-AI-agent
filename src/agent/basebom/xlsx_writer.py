"""서식보존 New BOM xlsx 출력 (Master compact_oven_processor.write_to_master_excel 패턴 이식).

원본 base BOM을 openpyxl로 리로드해 **원본 셀/병합/스타일을 일절 건드리지 않고**
우측에 주석 컬럼(New P/No·변경점·변경사유·출처)을 가산하고, 변경 행(MODIFY/DELETE)에
주석을 기입, ADD 행은 하단에 append. 기존 셀을 덮어쓰지 않으므로 병합셀 충돌·행매핑
오기입 위험이 없다. ``validate_new_bom`` 게이트(출처 없음/NEW 임의품번 거부)를 통과해야 출력.
"""
from __future__ import annotations

import io

from src.agent.basebom.apply import NewBom, validate_new_bom
from src.agent.basebom.parser import BaseBom

__all__ = ["write_newbom_xlsx"]

_ANNOT_HEADERS = ["New P/No", "변경점", "변경사유", "출처(SRC)"]


def write_newbom_xlsx(bom: BaseBom, nb: NewBom, *, header_excel_row: int = 1) -> bytes:
    """BaseBom 원본 + NewBom → 서식보존 xlsx bytes.

    Raises:
        ValueError: validate_new_bom 위반(출처 없음 / NEW 임의품번).
        RuntimeError: 원본 bytes 없음.
    """
    viol = validate_new_bom(nb)
    if viol:
        raise ValueError("New BOM 검증 위반: " + "; ".join(viol))
    if not bom.source_bytes:
        raise RuntimeError("서식보존 출력에는 원본 base BOM bytes가 필요합니다(parse_base_bom).")

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(bom.source_bytes), data_only=False)
    ws = (
        wb[bom.sheet_name]
        if (bom.sheet_name and bom.sheet_name in wb.sheetnames)
        else wb[wb.sheetnames[0]]
    )

    c0 = ws.max_column + 1  # 주석 컬럼 시작(원본 우측, 기존 셀 불침범)
    for k, h in enumerate(_ANNOT_HEADERS):
        ws.cell(row=header_excel_row, column=c0 + k, value=h)

    def _annot(excel_row: int, new_pno: str, status: str, reason: str, src: str) -> None:
        for k, v in enumerate([new_pno, status, reason, src]):
            if v:
                ws.cell(row=excel_row, column=c0 + k, value=v)

    for r in nb.rows:
        if r.status == "KEEP" or r.excel_row is None:
            continue
        new_pno = r.part_no_new if r.status == "MODIFY" else ("(삭제)" if r.status == "DELETE" else "")
        _annot(r.excel_row, new_pno, r.status, r.reason, r.src)

    # ADD 행(원본에 없던 신규) — 하단 append. 미발번 신규는 <발번대기>(apply가 강제).
    nxt = ws.max_row + 1
    for r in nb.rows:
        if r.status != "ADD":
            continue
        ws.cell(row=nxt, column=1, value=r.part_name or "")  # 부품명(근사 위치)
        _annot(nxt, r.part_no_new, "ADD", r.reason, r.src)
        nxt += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
