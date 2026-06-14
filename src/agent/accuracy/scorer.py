"""통합 master xlsx 정확도 채점기 (produced vs oracle 정답지).

목적
----
파이프라인이 산출한 "통합 개발부품 Master" xlsx를 정답지
(`통합 개발부품Master Compact v1.1.xlsx`, sheet "Master")와 비교해
**변경 유니버스**(분류 ∈ {New, Change, Delete}) 한정으로 정확도를 측정한다.

핵심 설계 결정
--------------
1. 분모는 변경 유니버스에만 둔다. Common(공통 carryover) 166행으로 정확도를
   희석하지 않는다. Common 회수 여부는 ``common_carryover_coverage``로 별도 보고.
2. base_pno 단독 키 금지 — v1.1에서 한 base_pno가 최대 5회 반복(14 collision
   그룹, 32행). produced/oracle in-scope 행을 1:1 **이분 매칭**으로 정렬하고
   blended 유사도(base/new/이름/분류 가중합)를 최대화한다. scipy
   ``linear_sum_assignment``가 있으면 최적, 없으면 greedy fallback.
3. 게이팅(헤드라인): base_pno, new_pno(정답에 실제 new P/No가 있는 행만),
   classification(정규화 후) exact-match.
4. 진단(soft, 비게이팅): part_name 토큰셋 비율, changing_point 3-class
   {add/change/delete}(분류에서 파생 — v1.1 변경점은 ~17개 자유 한글값이라
   exact/fuzzy 모두 실패), changing_reason fuzzy(difflib; 임베딩 옵션).
5. 정규화: NFC + casefold + strip. 분류 {신규↔New, 변경↔Change, 기존↔Common,
   삭제↔Delete}. new_pno '없음'/'-'/'←'/'X'/'' = 변경 없음(=base). 분류 'X'=Delete.
6. 제외 컬럼: qty_new, supplier, FMEA/HSMS/grade(v1.1 자체 미수록).

순수 함수 + dataclass. DB/LLM/네트워크 의존 없음.
"""

from __future__ import annotations

import difflib
import math
import os
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Optional, Sequence

# ──────────────────────────────────────────────────────────────────────────
# 정규화 헬퍼
# ──────────────────────────────────────────────────────────────────────────

# new_pno가 "변경 없음"(= base와 동일)을 의미하는 마커들.
_NOCHANGE_NEW_PNO = {"", "-", "←", "없음", "x"}

# 분류 정규화: 한국어/영어/대소문자 → canonical {New, Change, Common, Delete}.
_CLASSIFICATION_CANON: dict[str, str] = {
    "new": "New",
    "신규": "New",
    "change": "Change",
    "변경": "Change",
    "common": "Common",
    "기존": "Common",
    "공통": "Common",
    "delete": "Delete",
    "삭제": "Delete",
    "x": "Delete",  # 'X' 분류 마커 = Delete
}

# 변경 유니버스(채점 분모). Common은 carryover로 별도 집계.
_IN_SCOPE = {"New", "Change", "Delete"}

# changing_point 3-class 파생 매핑(분류 기반).
_CP_3CLASS = {"New": "add", "Change": "change", "Delete": "delete"}

_TOKEN_RE = re.compile(r"[^\w]+", re.UNICODE)


def _norm_text(value: Any) -> str:
    """NFC 정규화 + casefold + strip. None/빈값 → ''."""
    if value is None:
        return ""
    s = unicodedata.normalize("NFC", str(value))
    return s.strip().casefold()


def _norm_pno(value: Any) -> str:
    """P/No 정규화 — 텍스트 정규화에 더해 내부 공백 제거."""
    s = _norm_text(value)
    return s.replace(" ", "")


def _norm_classification(value: Any) -> str:
    """분류 → canonical {New, Change, Common, Delete} 또는 원문(미매핑 시)."""
    s = _norm_text(value)
    if not s:
        return ""
    return _CLASSIFICATION_CANON.get(s, s)


def _is_real_new_pno(raw_new: Any, base_norm: str) -> bool:
    """new_pno가 '실제 변경된 새 품번'인지 — 마커/공백/base와 동일이면 False."""
    s = _norm_pno(raw_new)
    if s in _NOCHANGE_NEW_PNO:
        return False
    if s == base_norm:  # base와 동일 표기 = 변경 없음
        return False
    return True


def _effective_new_pno(raw_new: Any, base_norm: str) -> str:
    """비교용 new_pno 값. '변경 없음' 마커류는 base와 동일로 취급."""
    if _is_real_new_pno(raw_new, base_norm):
        return _norm_pno(raw_new)
    return base_norm


def _tokens(value: Any) -> set[str]:
    s = _norm_text(value)
    if not s:
        return set()
    return {t for t in _TOKEN_RE.split(s) if t}


def _token_set_ratio(a: Any, b: Any) -> float:
    """토큰셋 자카드 유사도 [0,1]. 둘 다 비면 1.0, 한쪽만 비면 0.0."""
    ta, tb = _tokens(a), _tokens(b)
    if not ta and not tb:
        return 1.0
    if not ta or not tb:
        return 0.0
    inter = len(ta & tb)
    union = len(ta | tb)
    return inter / union if union else 0.0


def _seq_ratio(a: Any, b: Any) -> float:
    """difflib SequenceMatcher 비율 [0,1]."""
    sa, sb = _norm_text(a), _norm_text(b)
    if not sa and not sb:
        return 1.0
    if not sa or not sb:
        return 0.0
    return difflib.SequenceMatcher(None, sa, sb).ratio()


# ──────────────────────────────────────────────────────────────────────────
# 행 모델
# ──────────────────────────────────────────────────────────────────────────

# load_master_rows가 반환하는 정규화 키.
_ROW_KEYS = (
    "base_pno",
    "new_pno",
    "part_name",
    "bom_level",
    "part_type",
    "qty_base",
    "changing_point",
    "changing_reason",
    "classification",
)


@dataclass
class MasterRow:
    """Master 시트의 한 데이터 행(원시 셀값; 채점 시 정규화)."""

    base_pno: Any = ""
    new_pno: Any = ""
    part_name: Any = ""
    bom_level: Any = ""
    part_type: Any = ""
    qty_base: Any = ""
    changing_point: Any = ""
    changing_reason: Any = ""
    classification: Any = ""
    _source_row: Optional[int] = None  # 정답지 행 번호(있으면) — 에러 리포트용

    def as_dict(self) -> dict[str, Any]:
        return {k: getattr(self, k) for k in _ROW_KEYS}


def _coerce_row(item: Any) -> MasterRow:
    """dict 또는 MasterRow → MasterRow. 누락 키는 ''."""
    if isinstance(item, MasterRow):
        return item
    if isinstance(item, dict):
        kwargs = {k: item.get(k, "") for k in _ROW_KEYS}
        src = item.get("_source_row")
        return MasterRow(_source_row=src, **kwargs)
    raise TypeError(f"row must be dict or MasterRow, got {type(item)!r}")


# ──────────────────────────────────────────────────────────────────────────
# 헤더 탐지 + xlsx 로더
# ──────────────────────────────────────────────────────────────────────────

# 헤더 셀 텍스트 → 정규화 키 매핑(부분 일치, NFC+casefold 후).
# v1.1은 헤더가 2행(group + sub)로 쪼개짐. 'P/No' 그룹 + 'Base/New P/No' sub.
_HEADER_HINTS: list[tuple[str, str]] = [
    ("base p/no", "base_pno"),
    ("new p/no", "new_pno"),
    ("부품명", "part_name"),
    ("part name", "part_name"),
    ("class desc", "part_name"),
    ("bom\nlevel", "bom_level"),
    ("bom level", "bom_level"),
    ("part type", "part_type"),
    ("변경점", "changing_point"),
    ("changing point", "changing_point"),
    ("변경사유", "changing_reason"),
    ("changing reason", "changing_reason"),
    ("classification", "classification"),
    ("신규/변경", "classification"),
]


def _header_text(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFC", str(value)).strip().casefold()


def _row_header_keys(values: Sequence[Any]) -> dict[str, int]:
    """한 행에서 (정규화 키 → 컬럼 인덱스) 추출(부분 일치)."""
    found: dict[str, int] = {}
    for cidx, cell in enumerate(values):
        txt = _header_text(cell)
        if not txt:
            continue
        for hint, key in _HEADER_HINTS:
            if hint in txt and key not in found:
                found[key] = cidx
    return found


def _is_header_row(values: Sequence[Any]) -> bool:
    """행 자체에 Base/New P/No 또는 Part Name류 헤더가 있으면 True.

    v1.1은 헤더가 group행(부품명/변경점/분류…) + sub행(Base/New P/No)로
    쪼개져 있어 두 행 모두 이 조건을 만족한다 → load 시엔 '키 커버리지'로
    group행을 고른다(아래 _find_header_block).
    """
    joined = " | ".join(_header_text(v) for v in values)
    has_base = "base p/no" in joined
    has_new = "new p/no" in joined
    has_name = ("부품명" in joined) or ("part name" in joined) or ("class desc" in joined)
    return has_base or has_new or has_name


# 헤더 행 식별에 꼭 필요한 핵심 키(이 중 다수가 한 블록에 모여야 진짜 헤더).
_REQUIRED_HEADER_KEYS = {"base_pno", "new_pno", "part_name", "classification"}


def _find_header_block(rows: Sequence[Sequence[Any]]) -> tuple[int, int]:
    """헤더 블록(group행, 데이터 시작행) 탐지.

    각 행 i에 대해 (i, i+1) 2행 윈도우가 커버하는 정규화 키 수를 세고,
    핵심 키를 가장 많이 포함하는 윈도우의 **상단 group행**을 header_idx로
    고른다. v1.1처럼 split된 경우 group행(부품명·변경점·분류 포함)이 선택되고
    sub행(Base/New P/No)은 col map 병합에만 쓰인다.

    Returns:
        (header_idx, data_start_idx)  — 둘 다 0-based.
    """
    best_idx = -1
    best_score = -1
    for i in range(len(rows)):
        keys_i = set(_row_header_keys(rows[i]))
        keys_next = set(_row_header_keys(rows[i + 1])) if i + 1 < len(rows) else set()
        window = keys_i | keys_next
        # 핵심 키 커버리지를 우선, 동률이면 전체 키 수.
        score = len(window & _REQUIRED_HEADER_KEYS) * 100 + len(window)
        # group행이 되려면 그 행 자체가 sub-only(Base/New P/No뿐)가 아니어야 함.
        is_sub_only = keys_i and keys_i.issubset({"base_pno", "new_pno"})
        if is_sub_only:
            continue
        if window & _REQUIRED_HEADER_KEYS and score > best_score:
            best_score = score
            best_idx = i
    if best_idx < 0:
        # group행이 따로 없는 단일 헤더 레이아웃 — 첫 header_row 사용.
        single = next(
            (i for i, r in enumerate(rows) if _is_header_row(r)), None
        )
        if single is None:
            return -1, -1
        return single, single + 1
    # sub행(다음 행)이 Base/New P/No를 보강하면 데이터는 +2, 아니면 +1.
    next_keys = (
        set(_row_header_keys(rows[best_idx + 1]))
        if best_idx + 1 < len(rows)
        else set()
    )
    has_subheader = bool(next_keys & {"base_pno", "new_pno"})
    data_start = best_idx + (2 if has_subheader else 1)
    return best_idx, data_start


def _build_column_map(
    rows: Sequence[Sequence[Any]],
    header_idx: int,
) -> dict[str, int]:
    """group행 + 인접 sub행에서 (정규화 키 → 컬럼 인덱스) 병합 맵."""
    col_map: dict[str, int] = {}
    for ridx in (header_idx, header_idx + 1):
        if 0 <= ridx < len(rows):
            for key, cidx in _row_header_keys(rows[ridx]).items():
                col_map.setdefault(key, cidx)
    return col_map


def _detect_no_column(
    rows: Sequence[Sequence[Any]],
    header_idx: int,
) -> int:
    """'No.' 카운터 컬럼 인덱스 탐지. 못 찾으면 1(=col B) 기본."""
    scan = [header_idx]
    if header_idx + 1 < len(rows):
        scan.append(header_idx + 1)
    for ridx in scan:
        for cidx, cell in enumerate(rows[ridx]):
            if _header_text(cell) in ("no.", "no"):
                return cidx
    return 1  # default col B (0-based index 1)


def load_master_rows(path: str) -> list[dict]:
    """xlsx "Master" 시트 → 정규화 키 dict 리스트.

    헤더 행을 'Base P/No'/'New P/No'/'Part Name'(또는 한국어) 탐지로 찾고,
    'No.' 컬럼이 빈 행은 건너뛴다. produced/oracle 동일 v1.1 레이아웃이므로
    같은 로더로 둘 다 읽는다.

    Returns:
        per-row dict (keys: base_pno, new_pno, part_name, bom_level, part_type,
        qty_base, changing_point, changing_reason, classification, _source_row).
    """
    import openpyxl  # 지연 임포트(테스트는 xlsx 불필요)

    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    if "Master" not in wb.sheetnames:
        raise ValueError(f"'Master' sheet not found in {path!r}; sheets={wb.sheetnames}")
    ws = wb["Master"]

    # 전체를 행렬로(메모리 부담 없음 — ~243행).
    grid: list[list[Any]] = [
        [c.value for c in row] for row in ws.iter_rows()
    ]

    header_idx, data_start = _find_header_block(grid)
    if header_idx < 0:
        raise ValueError(f"header row (Base/New P/No) not found in {path!r}")

    col_map = _build_column_map(grid, header_idx)
    no_col = _detect_no_column(grid, header_idx)

    out: list[dict] = []
    for ridx in range(data_start, len(grid)):
        row = grid[ridx]

        def _get(key: str) -> Any:
            ci = col_map.get(key)
            if ci is None or ci >= len(row):
                return ""
            v = row[ci]
            return "" if v is None else v

        no_val = row[no_col] if no_col < len(row) else None
        if no_val is None or str(no_val).strip() == "":
            continue  # 'No.' 빈 행 = 데이터 아님

        out.append(
            {
                "base_pno": _get("base_pno"),
                "new_pno": _get("new_pno"),
                "part_name": _get("part_name"),
                "bom_level": _get("bom_level"),
                "part_type": _get("part_type"),
                "qty_base": _get("qty_base"),
                "changing_point": _get("changing_point"),
                "changing_reason": _get("changing_reason"),
                "classification": _get("classification"),
                "_source_row": ridx + 1,  # 1-based 시트 행 번호
            }
        )
    return out


# ──────────────────────────────────────────────────────────────────────────
# 이분 매칭(1:1)
# ──────────────────────────────────────────────────────────────────────────

# blended 유사도 가중치(합 = 1.0).
_W_BASE = 0.40
_W_NEW = 0.25
_W_NAME = 0.20
_W_CLS = 0.15


@dataclass
class _NormRow:
    """정규화된 in-scope 행(매칭/채점 단위)."""

    raw: MasterRow
    base: str
    new: str  # effective new (마커류는 base로 접힘)
    new_is_real: bool
    new_raw_norm: str
    name: str
    cls: str
    cp_class: str  # 3-class {add/change/delete}
    reason: str

    @classmethod
    def from_row(cls, row: MasterRow) -> "_NormRow":
        base = _norm_pno(row.base_pno)
        new_real = _is_real_new_pno(row.new_pno, base)
        return cls(
            raw=row,
            base=base,
            new=_effective_new_pno(row.new_pno, base),
            new_is_real=new_real,
            new_raw_norm=_norm_pno(row.new_pno),
            name=_norm_text(row.part_name),
            cls=_norm_classification(row.classification),
            cp_class=_CP_3CLASS.get(_norm_classification(row.classification), "?"),
            reason=str(row.changing_reason or ""),
        )


def _pair_similarity(p: _NormRow, o: _NormRow) -> float:
    """produced행 p vs oracle행 o의 blended 유사도 [0,1]."""
    base_sim = 1.0 if (p.base and p.base == o.base) else 0.0
    new_sim = 1.0 if p.new == o.new else 0.0
    name_sim = _token_set_ratio(p.name, o.name)
    cls_sim = 1.0 if (p.cls and p.cls == o.cls) else 0.0
    return (
        _W_BASE * base_sim
        + _W_NEW * new_sim
        + _W_NAME * name_sim
        + _W_CLS * cls_sim
    )


def _has_correspondence(p: _NormRow, o: _NormRow) -> bool:
    """매칭 수용 조건 — 실제 대응 신호가 있어야 쌍을 인정한다.

    base/new 품번 정확 일치 또는 부품명 토큰셋 유사 ≥0.5. **분류(cls)만 같은 약한
    신호로는 매칭하지 않는다** — 그러면 엉뚱한 행이 매칭돼 분류 점수가 과대계상된다
    (2026-06-13 적대적 리뷰 지적). 이 조건을 통과하지 못한 oracle 행은 unmatched로 남는다.
    """
    if p.base and p.base == o.base:
        return True
    if p.new_is_real and o.new_is_real and p.new == o.new:
        return True
    return _token_set_ratio(p.name, o.name) >= 0.5


def _match_bipartite(
    produced: list[_NormRow],
    oracle: list[_NormRow],
) -> tuple[list[tuple[int, int]], list[int], list[int]]:
    """produced/oracle in-scope 행 1:1 매칭(blended 유사도 최대화).

    scipy linear_sum_assignment 우선, 없으면 greedy fallback.
    유사도 0(완전 무관)인 쌍은 매칭에서 제외한다.

    Returns:
        (matched_pairs, unmatched_produced_idx, unmatched_oracle_idx)
        matched_pairs = [(prod_idx, oracle_idx), ...]
    """
    np_ = len(produced)
    no = len(oracle)
    if np_ == 0 or no == 0:
        return [], list(range(np_)), list(range(no))

    sim = [[_pair_similarity(p, o) for o in oracle] for p in produced]

    pairs: list[tuple[int, int]] = []
    try:
        import numpy as np
        from scipy.optimize import linear_sum_assignment

        cost = np.array(sim, dtype=float)
        # 비용 = -유사도(최대화 → 최소화).
        row_ind, col_ind = linear_sum_assignment(-cost)
        for r, c in zip(row_ind.tolist(), col_ind.tolist()):
            # 무관 쌍 제외 + 실제 대응 신호(base/new/이름) 필수(분류만으로 매칭 금지).
            if sim[r][c] > 0.0 and _has_correspondence(produced[r], oracle[c]):
                pairs.append((r, c))
    except Exception:  # noqa: BLE001 — scipy/numpy 없거나 실패 → greedy
        pairs = _match_greedy(sim, produced, oracle)

    used_p = {a for a, _ in pairs}
    used_o = {b for _, b in pairs}
    unmatched_p = [i for i in range(np_) if i not in used_p]
    unmatched_o = [i for i in range(no) if i not in used_o]
    return pairs, unmatched_p, unmatched_o


def _match_greedy(
    sim: list[list[float]],
    produced: list[_NormRow] | None = None,
    oracle: list[_NormRow] | None = None,
) -> list[tuple[int, int]]:
    """greedy 1:1 매칭 — 유사도 내림차순으로 미사용 쌍 선점.

    produced/oracle 주어지면 ``_has_correspondence`` 게이트 적용(분류만의 약한 매칭 방지).
    """
    cands: list[tuple[float, int, int]] = []
    for i, row in enumerate(sim):
        for j, s in enumerate(row):
            if s > 0.0 and (
                produced is None or oracle is None
                or _has_correspondence(produced[i], oracle[j])
            ):
                cands.append((s, i, j))
    cands.sort(key=lambda t: (-t[0], t[1], t[2]))
    used_p: set[int] = set()
    used_o: set[int] = set()
    pairs: list[tuple[int, int]] = []
    for s, i, j in cands:
        if i in used_p or j in used_o:
            continue
        used_p.add(i)
        used_o.add(j)
        pairs.append((i, j))
    return pairs


# ──────────────────────────────────────────────────────────────────────────
# changing_reason 유사도(difflib 기본, 임베딩 옵션)
# ──────────────────────────────────────────────────────────────────────────


def _cosine(a: Sequence[float], b: Sequence[float]) -> float:
    dot = sum(x * y for x, y in zip(a, b))
    na = math.sqrt(sum(x * x for x in a))
    nb = math.sqrt(sum(y * y for y in b))
    if na == 0.0 or nb == 0.0:
        return 0.0
    return dot / (na * nb)


def _reason_similarities(
    pairs_texts: list[tuple[str, str]],
) -> list[float]:
    """매칭 쌍별 changing_reason 유사도.

    ENABLE_EMBEDDING 설정 + src.embed.embedder.embed_texts 임포트 가능 시
    코사인 유사도, 아니면 difflib SequenceMatcher.
    임베딩 실패 시 difflib로 graceful fallback.
    """
    use_embed = os.environ.get("ENABLE_EMBEDDING", "").strip() not in ("", "0")
    if use_embed:
        try:
            from src.embed.embedder import embed_texts

            prod_texts = [a for a, _ in pairs_texts]
            ora_texts = [b for _, b in pairs_texts]
            all_texts = prod_texts + ora_texts
            vecs = embed_texts(all_texts)
            n = len(prod_texts)
            sims: list[float] = []
            for i in range(n):
                sims.append(_cosine(vecs[i], vecs[n + i]))
            return sims
        except Exception:  # noqa: BLE001 — 임베딩 불가 → difflib fallback
            pass
    return [_seq_ratio(a, b) for a, b in pairs_texts]


# ──────────────────────────────────────────────────────────────────────────
# 채점 코어
# ──────────────────────────────────────────────────────────────────────────

_NAME_MATCH_THRESHOLD = 0.85


@dataclass
class _Acc:
    """exact-match 누산기."""

    correct: int = 0
    total: int = 0

    def add(self, ok: bool) -> None:
        self.total += 1
        if ok:
            self.correct += 1

    def as_dict(self, key: str = "acc") -> dict[str, Any]:
        acc = (self.correct / self.total) if self.total else None
        return {key: acc, "matched": self.correct, "total": self.total}


def score_rows(
    produced_rows: Sequence[Any],
    oracle_rows: Sequence[Any],
) -> dict:
    """채점 코어 — 행 리스트(dict 또는 MasterRow)를 받아 결과 dict 반환.

    score()는 xlsx를 로드한 뒤 이 함수를 호출한다. 테스트는 이 함수를 직접 쓴다.
    """
    prod = [_coerce_row(r) for r in produced_rows]
    ora = [_coerce_row(r) for r in oracle_rows]

    prod_norm_all = [_NormRow.from_row(r) for r in prod]
    ora_norm_all = [_NormRow.from_row(r) for r in ora]

    # ── Common carryover coverage (별도 집계, 게이팅 분모에서 제외) ──
    ora_common = [r for r in ora_norm_all if r.cls == "Common"]
    prod_common = [r for r in prod_norm_all if r.cls == "Common"]
    common_cov = _common_carryover_coverage(prod_common, ora_common)

    # ── in-scope(변경 유니버스)로 스코프 ──
    prod_in = [r for r in prod_norm_all if r.cls in _IN_SCOPE]
    ora_in = [r for r in ora_norm_all if r.cls in _IN_SCOPE]

    pairs, unmatched_p, unmatched_o = _match_bipartite(prod_in, ora_in)

    # ── 게이팅 컬럼 ──
    base_acc = _Acc()
    new_acc = _Acc()
    cls_acc = _Acc()

    # ── 진단 컬럼 ──
    name_acc = _Acc()
    cp_acc = _Acc()
    reason_pairs_text: list[tuple[str, str]] = []

    errors: list[dict[str, Any]] = []

    for pi, oi in pairs:
        p = prod_in[pi]
        o = ora_in[oi]
        o_rownum = o.raw._source_row

        # base_pno (게이팅)
        base_ok = p.base == o.base
        base_acc.add(base_ok)
        if not base_ok:
            errors.append(
                _err("base_pno", o_rownum, "base_pno", o.base, p.base)
            )

        # new_pno (게이팅) — 정답에 '실제 new P/No'가 있는 행만 분모에 포함
        if o.new_is_real:
            new_ok = p.new_raw_norm == o.new_raw_norm
            new_acc.add(new_ok)
            if not new_ok:
                errors.append(
                    _err("new_pno", o_rownum, "new_pno", o.new_raw_norm, p.new_raw_norm)
                )

        # classification (게이팅, 정규화 후)
        cls_ok = p.cls == o.cls
        cls_acc.add(cls_ok)
        if not cls_ok:
            errors.append(
                _err("classification", o_rownum, "classification", o.cls, p.cls)
            )

        # part_name (진단) — 토큰셋 비율 >= 0.85
        name_ratio = _token_set_ratio(p.name, o.name)
        name_acc.add(name_ratio >= _NAME_MATCH_THRESHOLD)

        # changing_point 3-class (진단)
        cp_acc.add(p.cp_class == o.cp_class)

        # changing_reason (진단) — 매칭 후 일괄 계산
        reason_pairs_text.append((p.reason, o.reason))

    # 매칭 안 된 정답 행 = recall loss
    for oi in unmatched_o:
        o = ora_in[oi]
        errors.append(
            _err("unmatched_oracle", o.raw._source_row, None, _row_label(o), None)
        )
    # 매칭 안 된 산출 행 = precision loss
    for pi in unmatched_p:
        p = prod_in[pi]
        errors.append(
            _err("unmatched_produced", p.raw._source_row, None, None, _row_label(p))
        )

    reason_sims = _reason_similarities(reason_pairs_text)
    reason_mean = (sum(reason_sims) / len(reason_sims)) if reason_sims else None

    # ── overall gating accuracy: base/new/cls 가중 평균(샘플 수 기준) ──
    gate_correct = base_acc.correct + new_acc.correct + cls_acc.correct
    gate_total = base_acc.total + new_acc.total + cls_acc.total
    overall_gating = (gate_correct / gate_total) if gate_total else None

    n_matched = len(pairs)
    n_match_denom = max(len(prod_in), len(ora_in))
    match_rate = (n_matched / n_match_denom) if n_match_denom else 1.0

    return {
        "scope": {
            "oracle_inscope_rows": len(ora_in),
            "produced_inscope_rows": len(prod_in),
        },
        "row_matching": {
            "matched": n_matched,
            "unmatched_produced": len(unmatched_p),
            "unmatched_oracle": len(unmatched_o),
            "match_rate": match_rate,
        },
        "gating": {
            "base_pno": base_acc.as_dict(),
            "new_pno": new_acc.as_dict(),
            "classification": cls_acc.as_dict(),
            "overall_gating_acc": overall_gating,
        },
        "diagnostic": {
            "part_name": name_acc.as_dict("match_rate"),
            "changing_point_3class": cp_acc.as_dict("agreement"),
            "changing_reason": {
                "mean_similarity": reason_mean,
                "method": _reason_method(),
                "n": len(reason_sims),
            },
        },
        "common_carryover_coverage": common_cov,
        "errors": errors,
    }


def _reason_method() -> str:
    use_embed = os.environ.get("ENABLE_EMBEDDING", "").strip() not in ("", "0")
    if use_embed:
        try:
            import importlib.util

            if importlib.util.find_spec("src.embed.embedder") is not None:
                return "embedding_cosine"
        except Exception:  # noqa: BLE001
            pass
    return "difflib"


def _common_carryover_coverage(
    prod_common: list[_NormRow],
    ora_common: list[_NormRow],
) -> dict[str, Any]:
    """Common 행 회수율 — produced가 정답 Common base_pno들을 (multiset) 회수했나.

    게이팅과 무관(별도 보고). base_pno multiset 교집합 / oracle Common 수.
    """
    from collections import Counter

    ora_bases = Counter(r.base for r in ora_common if r.base)
    prod_bases = Counter(r.base for r in prod_common if r.base)
    covered = sum((ora_bases & prod_bases).values())
    total = sum(ora_bases.values())
    coverage = (covered / total) if total else None
    return {
        "oracle_common_rows": len(ora_common),
        "produced_common_rows": len(prod_common),
        "covered": covered,
        "coverage": coverage,
    }


def _row_label(r: _NormRow) -> str:
    return f"{r.base}|{r.new}|{r.cls}"


def _err(
    type_: str,
    oracle_row: Any,
    column: Optional[str],
    expected: Any,
    produced: Any,
) -> dict[str, Any]:
    return {
        "type": type_,
        "oracle_row": oracle_row,
        "column": column,
        "expected": expected,
        "produced": produced,
    }


# ──────────────────────────────────────────────────────────────────────────
# 공개 진입점
# ──────────────────────────────────────────────────────────────────────────


def score(produced_path: str, oracle_path: str) -> dict:
    """produced xlsx vs oracle 정답지 xlsx 비교 → 채점 결과 dict.

    Args:
        produced_path: 파이프라인 산출 통합 master xlsx 경로.
        oracle_path:   정답지(`통합 개발부품Master Compact v1.1.xlsx`) 경로.

    Returns:
        score_rows()의 결과 dict.
    """
    produced_rows = load_master_rows(produced_path)
    oracle_rows = load_master_rows(oracle_path)
    return score_rows(produced_rows, oracle_rows)
