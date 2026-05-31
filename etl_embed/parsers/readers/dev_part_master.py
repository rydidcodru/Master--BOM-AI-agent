"""Group A: 통합 개발부품 Master 양식 파서.

대상 파일 (5개):
  - 동유럽향, 싱가포르향 (v1: R7-8 이중헤더)
  - 모리셔스향, 북유럽, 유럽 Steam Heater (v2: R8 단일헤더)

핵심 책임:
  1. 시트별 양식 변형(v1/v2) 자동 판별
  2. 상단 폼 메타(Base Model, New Model, Event 등) 추출
  3. 헤더 행에서 컬럼 위치 매핑
  4. 데이터 행을 dict 리스트로 변환
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..utils import (
    clean_cell, clean_header, is_empty_row,
    parse_bom_depth, parse_int, normalize_yn,
)


# 양식 변형 식별 (시트의 헤더 위치로 판별)
FORM_V1_SHEETS_HINT = {  # 동유럽/싱가포르 — R7=메인 헤더, R8=서브 헤더(Base/New)
    "Good-1 BK", "Good-2 BK", "Goood-2 BK", "Good-1 STS", "Good-2 STS",
    "Master(Best)", "Master(Better)", "Master(Good)",
}

# 양식 v1 컬럼 매핑 (R7 메인 헤더 텍스트 → DB 필드)
# 헤더는 줄바꿈/공백 정규화 후 매칭 (clean_header 적용 후 prefix 매칭)
V1_HEADER_MAP = {
    "No.": "_no",
    "BOM Level": "bom_level_raw",
    "Part Type": "part_type",
    "P/No": "_pno_group",        # R8에서 Base/New로 갈라짐
    "부품명": "part_name",
    "Class Desc.": "part_name",
    "Q'ty": "_qty_group",        # R8에서 Base/New로 갈라짐
    "변경점": "change_point_raw",
    "변경사유": "change_reason_raw",
    "양산처": "supplier",
    "Supplier": "supplier",
    "신규/변경": "classification",
    "Classification": "classification",
    "금형": "mold_dev",
    "Mold Dev": "mold_dev",
    "사내": "in_house",
    "In-house": "in_house",
}

# 양식 v2 컬럼 매핑 (모리셔스/북유럽/유럽 — R8 단일헤더)
V2_HEADER_MAP = {
    "No.": "_no",
    "부품 개발 등록": "_dev_register",
    "BOM Level": "bom_level_raw",
    "Part Type": "part_type",
    "금형 개발": "mold_dev",
    "사내 제작 or 직거래": "in_house",
    "Base P/No": "part_no_base",
    "New P/No": "part_no_new",
    "Class Desc.": "part_name",
    "Class Desc.(Part Name)": "part_name",
    "Fig.": "fig_ref",
    "변경점": "change_point_raw",
    "구분": "classification",
    "변경사유": "change_reason_raw",
    "Quanty": "qty_new",
    "Quantity": "qty_new",
    "Q'ty": "qty_new",
    "설계자": "designer",
}


HEADER_SCAN_MAX_ROW = 20
HEADER_SCAN_MAX_COL = 80

FIELD_ALIASES = {
    "no": ("No.", "No"),
    "bom_level_raw": ("BOM Level", "Level", "Lvl"),
    "part_type": ("Part Type", "분과", "CMDT", "CMDT-2"),
    "part_no_base": ("Base P/No", "Base Part No", "Base P/no", "Base P/No."),
    "part_no_new": (
        "New P/No", "New Part No", "New P/no", "New P/No.",
        "P/no.", "P/No.", "Part No", "Part No.", "Part No", "P/no",
    ),
    "part_name": (
        "Class Desc.(Part Name)", "Class Desc.", "Part Name", "Desc.",
        "Description", "부품명", "Part",
    ),
    "qty_base": ("Base Q'ty", "Base Qty", "Base Quantity"),
    "qty_new": ("New Q'ty", "New Qty", "New Quantity", "Q'ty", "Qty", "QTY", "Quantity", "Quanty"),
    "change_point_raw": ("Changing Point", "Changing Point.", "변경점", "변경 내역", "변경내역"),
    "change_reason_raw": ("Changing Reason", "Changing Reason.", "변경사유", "변경 사유"),
    "supplier": ("Supplier", "양산처", "Supplier Code"),
    "classification": ("Classification", "신규/변경", "구분", "신규"),
    "mold_dev": ("Mold Dev", "Mold Dev.", "금형 개발", "금형", "금형 개발 /수정"),
    "in_house": ("In-house", "In-house.", "사내 제작 or 직거래", "사내 입고", "사내"),
    "designer": ("설계자", "설계"),
    "fig_ref": ("Fig.", "Fig", "Picture"),
    "part_grade": ("부품 등급", "Grade"),
    "remark": ("Remark", "비고", "비 고"),
    "module": ("Module",),
    "buyer": ("구매",),
    "introduced": ("도입",),
}

EXTRA_FIELDS = {"module", "buyer", "introduced"}

HEADER_SCORE_FIELDS = {
    "bom_level_raw", "part_no_base", "part_no_new", "part_name",
    "classification", "change_point_raw", "change_reason_raw",
}

DYNAMIC_SHEET_EXCLUDE_KEYWORDS = (
    "history", "summary", "dqms", "hsms", "ssoid", "fcm", "partdms", "pra",
    "시험", "인정시험", "부품등급", "등급심의", "비교단가", "bom compare",
)


@dataclass
class FormMeta:
    """시트 상단에서 추출한 폼 메타데이터."""
    base_model: str | None = None
    base_grade: str | None = None
    new_model: str | None = None
    new_grade: str | None = None
    event: str | None = None
    buyer: str | None = None
    brand: str | None = None
    set_part_no: str | None = None
    production_date: str | None = None


@dataclass
class ParsedRow:
    """파싱된 1개 데이터 행 (DB INSERT용)."""
    sheet_name: str
    source_row: int
    form_id: str
    meta: FormMeta
    data: dict[str, Any] = field(default_factory=dict)
    extra_fields: dict[str, Any] = field(default_factory=dict)


@dataclass
class HeaderLayout:
    """Detected header layout for a sheet."""
    form_id: str
    header_row: int
    data_start_row: int
    col_idx: dict[str, int]
    requires_part_no: bool = True


# ===== Form variant detection =====

def detect_form_variant(ws: Worksheet) -> str | None:
    """시트가 v1(이중헤더) 인지 v2(단일헤더) 인지 판별.

    openpyxl row 인덱스 기준:
      - v1: row 8 = 메인 헤더 ("BOM Level"), row 9 = 서브 헤더 ("Base P/No")
      - v2: row 9 = 단일 헤더 ("BOM Level"과 "Base P/No" 동시 등장)
    """
    try:
        rows = [list(ws.iter_rows(min_row=r, max_row=r, max_col=30, values_only=True))[0]
                for r in (8, 9, 10)]
    except IndexError:
        return None
    r8_norm = [clean_header(c) or "" for c in rows[0]]
    r9_norm = [clean_header(c) or "" for c in rows[1]]

    def has(arr, text):
        return any(text in (c or "") for c in arr)

    # v1: row 8에 'BOM Level' 메인, row 9에 'Base P/No' 서브
    if has(r8_norm, "BOM Level") and has(r9_norm, "Base P/No"):
        return "dev_part_master_v1"
    # v2: row 9에 'BOM Level'과 'Base P/No' 동시 등장
    if has(r9_norm, "BOM Level") and has(r9_norm, "Base P/No"):
        return "dev_part_master_v2"
    return None


def detect_header_layout(ws: Worksheet) -> HeaderLayout | None:
    """Detect supported dev-part layouts without relying on fixed row 8/9."""
    legacy_form = detect_form_variant(ws)
    if legacy_form == "dev_part_master_v1":
        return HeaderLayout(
            form_id=legacy_form,
            header_row=8,
            data_start_row=10,
            col_idx=build_column_index_v1(ws),
        )
    if legacy_form == "dev_part_master_v2":
        return HeaderLayout(
            form_id=legacy_form,
            header_row=9,
            data_start_row=10,
            col_idx=build_column_index_v2(ws),
        )

    if _is_dynamic_sheet_excluded(ws.title):
        return None

    best: tuple[int, int, dict[str, int]] | None = None
    max_row = min(ws.max_row or HEADER_SCAN_MAX_ROW, HEADER_SCAN_MAX_ROW)
    max_col = min(ws.max_column or HEADER_SCAN_MAX_COL, HEADER_SCAN_MAX_COL)
    for row_num in range(1, max_row + 1):
        rows = list(ws.iter_rows(
            min_row=row_num,
            max_row=row_num,
            max_col=max_col,
            values_only=True,
        ))
        if not rows:
            continue
        col_idx = build_column_index_from_headers(rows[0])
        score = _score_column_index(col_idx)
        if score < 3:
            continue
        if best is None or score > best[0]:
            best = (score, row_num, col_idx)

    if best is None:
        return None

    _score, header_row, col_idx = best
    requires_part_no = "part_no_new" in col_idx or "part_no_base" in col_idx
    return HeaderLayout(
        form_id="dev_part_master_dynamic",
        header_row=header_row,
        data_start_row=header_row + 1,
        col_idx=col_idx,
        requires_part_no=requires_part_no,
    )


def _is_dynamic_sheet_excluded(sheet_name: str) -> bool:
    key = sheet_name.strip().lower()
    return any(keyword in key for keyword in DYNAMIC_SHEET_EXCLUDE_KEYWORDS)


def build_column_index_from_headers(headers: tuple) -> dict[str, int]:
    """Map a single detected header row into canonical DB/extra fields."""
    idx: dict[str, int] = {}
    for col_i, raw_header in enumerate(headers):
        header = clean_header(raw_header)
        if not header:
            continue
        field_name = _match_header_alias(header)
        if not field_name or field_name in idx:
            continue
        idx[field_name] = col_i
    return idx


def _score_column_index(col_idx: dict[str, int]) -> int:
    score = sum(2 for field in HEADER_SCORE_FIELDS if field in col_idx)
    if "part_name" in col_idx:
        score += 2
    if "part_no_new" in col_idx or "part_no_base" in col_idx:
        score += 2
    if "change_point_raw" in col_idx and "change_reason_raw" in col_idx:
        score += 2
    return score


def _match_header_alias(header: str) -> str | None:
    header_key = _header_key(header)
    for field_name, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            alias_key = _header_key(alias)
            if header_key == alias_key:
                return field_name
    for field_name, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            alias_key = _header_key(alias)
            if len(alias_key) >= 4 and alias_key in header_key:
                return field_name
    return None


def _header_key(value: Any) -> str:
    text = clean_header(value) or ""
    text = text.lower()
    return re.sub(r"[^0-9a-z가-힣]+", "", text)


# ===== Metadata extraction =====

def extract_form_meta_v1(ws: Worksheet) -> FormMeta:
    """v1 양식: row 2(Base Model), row 3(New Model), row 4(Event).
    col 1(B열) 라벨, col 5(F열) 값.
    """
    meta = FormMeta()
    for r in range(2, 8):
        rows = list(ws.iter_rows(min_row=r, max_row=r, max_col=10, values_only=True))
        if not rows:
            continue
        row = rows[0]
        label = clean_header(row[1]) if len(row) > 1 else None
        value = clean_cell(row[5]) if len(row) > 5 else None
        if not label:
            continue
        if "Base Model" in label:
            base_model, base_grade = _split_model_grade(value)
            meta.base_model = base_model
            meta.base_grade = base_grade
        elif "New Model" in label:
            new_model, new_grade = _split_model_grade(value)
            meta.new_model = new_model
            meta.new_grade = new_grade
        elif "Event" in label:
            meta.event = str(value) if value is not None else None
    return meta


def extract_form_meta_v2(ws: Worksheet) -> FormMeta:
    """v2 양식: row 2-7에 폼 메타.
    col 7=라벨, col 8=base 값 / col 10=라벨, col 11=new 값.
      row 2: '모델명(등급)'  row 3: 'Buyer명'  row 4: 'Brand'
      row 5: 'Set P/No.'  row 6: '양산 일자'  row 7: 'Event' (col 9 라벨, col 10 값)
    """
    meta = FormMeta()
    for r in range(2, 8):
        rows = list(ws.iter_rows(min_row=r, max_row=r, max_col=14, values_only=True))
        if not rows:
            continue
        row = rows[0]
        label_base = clean_header(row[7]) if len(row) > 7 else None
        value_base = clean_cell(row[8]) if len(row) > 8 else None
        label_new = clean_header(row[10]) if len(row) > 10 else None
        value_new = clean_cell(row[11]) if len(row) > 11 else None

        if label_base and "모델명" in label_base:
            meta.base_model, meta.base_grade = _split_model_grade(value_base)
        if label_new and "모델명" in label_new:
            meta.new_model, meta.new_grade = _split_model_grade(value_new)
        if label_base and "Buyer" in label_base:
            meta.buyer = str(value_base) if value_base else None
        if label_base and "Brand" in label_base:
            meta.brand = str(value_base) if value_base else None
        if label_base and "Set P/No" in label_base:
            meta.set_part_no = str(value_base) if value_base else None
        if label_base and "양산 일자" in label_base:
            meta.production_date = str(value_base) if value_base else None

        # Event는 row 7 col 9='Event' label, col 10=value
        evt_label = clean_header(row[9]) if len(row) > 9 else None
        evt_value = clean_cell(row[10]) if len(row) > 10 else None
        if evt_label == "Event" and evt_value:
            meta.event = str(evt_value)
    return meta


def _split_model_grade(raw: Any) -> tuple[str | None, str | None]:
    """'WSED7613B / D' → ('WSED7613B', 'D')
       'WSED7613S / Cb (사우디)' → ('WSED7613S', 'Cb')
       'WS7D7632WB.ABKQEUE' → ('WS7D7632WB.ABKQEUE', None)
    """
    if raw is None:
        return None, None
    s = clean_cell(raw)
    if s is None:
        return None, None
    s = str(s)
    if "/" in s:
        parts = s.split("/", 1)
        model = parts[0].strip()
        grade = re.sub(r"\(.+?\)", "", parts[1]).strip()
        return model or None, grade or None
    return s, None


# ===== Header parsing =====

def build_column_index_v1(ws: Worksheet) -> dict[str, int]:
    """v1: row 8(메인)과 row 9(서브)를 결합해서 컬럼 인덱스 매핑.
    메인의 'P/No'는 서브에서 'Base P/No'/'New P/No'로 갈라짐.
    메인의 "Q'ty"는 서브에서 'Base'/'New'로 갈라짐.

    Returns: {DB field name: 0-based column index}
    """
    r8_rows = list(ws.iter_rows(min_row=8, max_row=8, max_col=30, values_only=True))
    if not r8_rows:
        return {}
    r8 = r8_rows[0]

    idx: dict[str, int] = {}
    for col_i, h in enumerate(r8):
        header = clean_header(h)
        if not header:
            continue
        for prefix, field_name in V1_HEADER_MAP.items():
            if header.startswith(prefix):
                if field_name == "_pno_group":
                    idx["part_no_base"] = col_i
                    idx["part_no_new"] = col_i + 1
                elif field_name == "_qty_group":
                    idx["qty_base"] = col_i
                    idx["qty_new"] = col_i + 1
                elif field_name == "_no":
                    pass
                else:
                    idx[field_name] = col_i
                break
    return idx


def build_column_index_v2(ws: Worksheet) -> dict[str, int]:
    """v2: row 9 단일헤더."""
    rows = list(ws.iter_rows(min_row=9, max_row=9, max_col=30, values_only=True))
    if not rows:
        return {}
    r9 = rows[0]
    idx: dict[str, int] = {}
    for col_i, h in enumerate(r9):
        header = clean_header(h)
        if not header:
            continue
        for prefix, field_name in V2_HEADER_MAP.items():
            if header.startswith(prefix):
                if field_name.startswith("_"):
                    continue
                idx[field_name] = col_i
                break
    return idx


# ===== Row parsing =====

def parse_data_row(row: tuple, col_idx: dict[str, int], known_fields: set[str]) -> tuple[dict, dict]:
    """1개 행을 (data_dict, extra_fields_dict)로 분해.

    known_fields에 있는 컬럼명은 정형 컬럼 (data_dict)으로,
    그 외는 모두 extra_fields (JSON)로.
    """
    data: dict[str, Any] = {}
    extra: dict[str, Any] = {}
    for field_name, col_i in col_idx.items():
        if col_i >= len(row):
            continue
        raw = row[col_i]
        if field_name in {"qty_base", "qty_new"}:
            value = parse_int(raw)
        elif field_name in {"mold_dev", "in_house"}:
            value = normalize_yn(raw)
        elif field_name == "bom_level_raw":
            value = clean_cell(raw)
            if value is not None:
                value = str(value)
        else:
            value = clean_cell(raw)
            if value is not None and not isinstance(value, (int, float)):
                value = str(value)

        if value is None:
            continue
        if field_name in known_fields:
            data[field_name] = value
        else:
            extra[field_name] = value
    return data, extra


# ===== Top-level reader =====

# dev_part_master 테이블의 정형 컬럼 (extra_fields에 들어가지 않을 것들)
DPM_STRUCTURED_FIELDS = {
    "bom_level_raw", "part_type", "part_no_base", "part_no_new", "part_name",
    "qty_base", "qty_new", "change_point_raw", "change_reason_raw",
    "supplier", "classification", "mold_dev", "in_house",
    "designer", "fig_ref", "part_grade", "remark",
}


import re as _re

# 부품번호 정규식 (예: AEC74157606, MFL71927538, EAD63668414)
_PART_NO_PATTERN = _re.compile(r"^[A-Z]{2,5}\d{5,12}$")


def _looks_like_part_no(value) -> bool:
    """문자열이 부품번호처럼 생겼는지."""
    if value is None:
        return False
    s = str(value).strip()
    return bool(_PART_NO_PATTERN.match(s))


def _normalize_classification(value) -> str | None:
    """다양한 표기 변형을 표준화.
    'Chainging', 'Changing', 'Change' → 'Changing'
    'NEW', 'New', 'new' → 'NEW'
    'Common', 'common' → 'Common'
    'Delet', 'Delete' → 'Delete'
    """
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    s_lower = s.lower()
    if s_lower in ("new",):
        return "NEW"
    if s_lower in ("chainging", "changing", "change"):
        return "Changing"
    if s_lower in ("common",):
        return "Common"
    if s_lower in ("delet", "delete"):
        return "Delete"
    return s  # 그 외는 원본 유지


def read_sheet(ws: Worksheet, file_path: Path) -> list[ParsedRow]:
    """시트 1개를 파싱해서 ParsedRow 리스트 반환.
    지원하지 않는 양식이면 빈 리스트.
    """
    layout = detect_header_layout(ws)
    if layout is None:
        return []

    form_id = layout.form_id
    if form_id == "dev_part_master_v1":
        meta = extract_form_meta_v1(ws)
    else:
        meta = extract_form_meta_v2(ws)

    rows_out: list[ParsedRow] = []
    for r_idx, row in enumerate(
        ws.iter_rows(
            min_row=layout.data_start_row,
            max_col=HEADER_SCAN_MAX_COL,
            values_only=True,
        ),
        start=layout.data_start_row,
    ):
        if is_empty_row(row):
            continue
        data, extra = parse_data_row(row, layout.col_idx, DPM_STRUCTURED_FIELDS)
        # part_name이 없으면 데이터 행이 아닌 것으로 간주 (footer, summary 등)
        if not data.get("part_name"):
            continue
        # 기본은 부품번호를 요구한다. 단, Part/변경내역 중심의 간이 변경표는 예외.
        has_part_no = (
            _looks_like_part_no(data.get("part_no_new"))
            or _looks_like_part_no(data.get("part_no_base"))
        )
        has_change_text = bool(data.get("change_point_raw") or data.get("change_reason_raw"))
        if layout.requires_part_no and not has_part_no:
            continue
        if not layout.requires_part_no and not has_change_text:
            continue
        # classification 정규화
        if "classification" in data:
            normalized = _normalize_classification(data["classification"])
            if normalized:
                data["classification"] = normalized
            else:
                data.pop("classification", None)
        # bom_depth 파생
        if "bom_level_raw" in data:
            depth = parse_bom_depth(data["bom_level_raw"])
            if depth is not None:
                data["bom_depth"] = depth

        rows_out.append(ParsedRow(
            sheet_name=ws.title,
            source_row=r_idx,
            form_id=form_id,
            meta=meta,
            data=data,
            extra_fields=extra,
        ))
    return rows_out


def read_file(file_path: str | Path) -> tuple[list[ParsedRow], list[str]]:
    """파일 1개를 파싱.
    Returns: (parsed_rows, skipped_sheets)
    """
    file_path = Path(file_path)
    wb = load_workbook(file_path, read_only=True, data_only=True)
    all_rows: list[ParsedRow] = []
    skipped: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        try:
            rows = read_sheet(ws, file_path)
        except Exception as e:  # noqa: BLE001
            skipped.append(f"{sheet_name}: ERROR {e}")
            continue
        if not rows:
            skipped.append(sheet_name)
        else:
            all_rows.extend(rows)
    wb.close()
    return all_rows, skipped
