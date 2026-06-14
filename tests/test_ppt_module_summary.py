"""PPT 'Module명/주요 변경점' 요약표 추출 테스트 (합성 deck, 네트워크 불필요)."""

from __future__ import annotations

import io

import pytest

pptx = pytest.importorskip("pptx")
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from src.agent.ppt.extractor import (  # noqa: E402
    change_items_from_extraction,
    extract_change_review_from_pptx_bytes,
)


def _deck() -> bytes:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    tf = s.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(9), Inches(0.6)).text_frame
    tf.text = "유첨 개발 변경점 상세"
    tbl = s.shapes.add_table(3, 3, Inches(0.3), Inches(1), Inches(9), Inches(2)).table
    for j, h in enumerate(["No.", "Module명", "주요 변경점"]):
        tbl.cell(0, j).text = h
    tbl.cell(1, 0).text = "1"
    tbl.cell(1, 1).text = "Cavity"
    tbl.cell(1, 2).text = "1. 제품 치수 변경에 따른 H 치수 137mm 감소 2. Cavity 조립 방식 변경"
    tbl.cell(2, 0).text = "2"
    tbl.cell(2, 1).text = "Insulator"
    tbl.cell(2, 2).text = "BO24 대비 변경점 없음"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_module_summary_extracted_to_change_items():
    ext = extract_change_review_from_pptx_bytes(_deck())
    items = change_items_from_extraction(ext)
    details = [it.change_detail for it in items]
    assert any("Cavity 조립 방식 변경" in d for d in details)   # 추출됨
    assert any("137mm" in d for d in details)                   # 번호 항목 분리
    assert not any("변경점 없음" in d for d in details)          # no-change 필터
    assert any(it.module == "Cavity" for it in items)           # 모듈 보존
    # 검색 텍스트는 변경내용만 (reason='정보 없음'은 제외)
    cavity = next(it for it in items if "Cavity 조립" in it.change_detail)
    assert "Cavity 조립 방식 변경" in cavity.search_text
    # 모듈요약표엔 변경사유 컬럼이 없지만, 인과절('A에 따른 B')에서 변경사유(A)를 추출.
    h_item = next(it for it in items if "137mm" in it.change_detail)
    assert h_item.change_reason == "제품 치수 변경"
    assert cavity.change_reason == "정보 없음"  # 인과어 없음 → 사유 미상(미생성)


def _deck_summary_plus_detail() -> bytes:
    """모듈요약표 + 유첨 상세표(변경사유 컬럼)가 공존하는 합성 덱."""
    prs = Presentation()
    # 슬라이드 1: 모듈요약표 (No.|Module명|주요 변경점) — Camera 라인은 사유 컬럼 없음.
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(9), Inches(0.5)).text_frame.text = "주요 변경점 요약"
    t1 = s1.shapes.add_table(2, 3, Inches(0.3), Inches(1), Inches(9), Inches(1.5)).table
    for j, h in enumerate(["No.", "Module명", "주요 변경점"]):
        t1.cell(0, j).text = h
    t1.cell(1, 0).text = "1"
    t1.cell(1, 1).text = "Door"
    t1.cell(1, 2).text = "Camera Module(LED) 장착 구조 반영"
    # 슬라이드 2: 유첨 상세표 (구분|No|Part|변경 내역|변경 사유|걱정점) — 사유 보유.
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(9), Inches(0.5)).text_frame.text = "유첨 개발 변경점 상세"
    t2 = s2.shapes.add_table(2, 6, Inches(0.3), Inches(1), Inches(9), Inches(1.5)).table
    for j, h in enumerate(["구분", "No", "Part", "변경 내역\n(변경 전 → 변경 후)", "변경 사유", "걱정점"]):
        t2.cell(0, j).text = h
    for j, v in enumerate(["기구", "1", "Cover, Camera", "Camera 모듈 장착으로 인한 부품 추가", "Door 카메라 적용", "온도"]):
        t2.cell(1, j).text = v
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_detail_table_suppresses_module_summary_fallback():
    """유첨 상세표(변경사유 보유)가 있으면 part/사유가 비는 모듈요약 fallback 항목은 빠진다."""
    ext = extract_change_review_from_pptx_bytes(_deck_summary_plus_detail())
    items = change_items_from_extraction(ext)
    details = [it.change_detail for it in items]
    # 모듈요약 fallback("Camera Module(LED) ... 장착 구조 반영")은 억제됨.
    assert not any("장착 구조 반영" in d for d in details)
    # 상세표의 사유 채워진 항목은 남고, 사유가 채워진다.
    cam = next(it for it in items if "Camera 모듈 장착" in it.change_detail)
    assert cam.change_reason == "Door 카메라 적용"


def test_split_cause_extracts_reason_from_causal_clause():
    from src.agent.ppt.extractor import _split_cause

    assert _split_cause("제품 치수 변경에 따른 H 치수 137mm 감소") == "제품 치수 변경"
    assert _split_cause("Conv Heater 삭제에 따른 인쇄 변경") == "Conv Heater 삭제"
    assert _split_cause("Door 무게 변경으로 soft closing hinge 수정") == "Door 무게 변경"
    assert _split_cause("원가 절감을 위해 재질 변경") == "원가 절감"
    assert _split_cause("Cavity 조립 방식 변경") == ""           # 인과어 없음 → 사유 미상
    assert _split_cause("MWO, Steam 기능 착탈 구조 반영") == ""
