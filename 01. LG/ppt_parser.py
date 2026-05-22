from __future__ import annotations
import re
from typing import Any


_UI_FORBIDDEN_TOKENS = ("http", "https", "www", "sharepoint", "onedrive", "teams", "drive", "path", "folder")


def _has_forbidden_ui_text(text: Any) -> bool:
    s = str(text or "").lower()
    return any(tok in s for tok in _UI_FORBIDDEN_TOKENS)


def _safe_ui_text(text: Any) -> str:
    s = str(text or "").strip()
    return "" if _has_forbidden_ui_text(s) else s


def _is_meaningful_part_no(text: Any) -> bool:
    s = re.sub(r"\s+", "", str(text or "").strip().upper())
    return bool(s and s not in {"TBD", "N/A", "NA", "NONE", "정보없음"})


_REMARK_BANNED_PATTERNS = (
    "낮은 확신도", "확신도", "low confidence", "confidence", "불확실", "추정",
)


def _sanitize_user_remark_text(text: Any) -> str:
    raw = str(text or "").strip()
    if not raw:
        return ""
    frags = re.split(r"\s*\|\s*", raw)
    kept = []
    for f in frags:
        s = str(f or "").strip()
        if not s:
            continue
        low = s.lower()
        if any(p in low for p in _REMARK_BANNED_PATTERNS):
            continue
        kept.append(s)
    return " | ".join(kept).strip(" |")


def _build_ui_meta_for_change_point(row: dict) -> tuple[str, str]:
    ctype = str(row.get("type") or "").strip()
    base_pno = str(row.get("base_part_no") or "").strip()
    new_pno = str(row.get("new_part_no") or "").strip()
    shared_src = str(row.get("shared_source") or "").strip()
    is_shared = bool(row.get("is_shared_part")) or bool(shared_src) or ("공용" in [str(t).strip() for t in (row.get("tags") or [])])
    concerns = row.get("concerns") or []
    concern0 = str(concerns[0]).strip() if concerns else ""
    desc = str(row.get("description") or "")

    ui_type = ""
    ui_remark = ""

    # =====================================
    # 공용품 우선 처리: 공용이면 채번 필요 로직 무시
    # =====================================
    if is_shared:
        if ctype == "NEW":
            ui_type = "신규(품번확정) · 공용"
            ui_remark = f"공용품 P/N: {new_pno}" if _is_meaningful_part_no(new_pno) else "공용품 | 품번 확인 필요"
        elif ctype == "Changing":
            ui_type = "변경 · 공용"
            if _is_meaningful_part_no(base_pno) and _is_meaningful_part_no(new_pno) and base_pno != new_pno:
                ui_remark = f"Base: {base_pno} → New: {new_pno} (공용)"
            else:
                ui_remark = "공용품"
        else:
            ui_type = "공용"
            ui_remark = "공용품"
    else:
        # =====================================
        # 공용 아닌 경우: 종전 로직
        # =====================================
        is_new_need = (
            ctype == "NEW"
            and (
                not _is_meaningful_part_no(new_pno)
                or "품번추가" in desc
                or "채번" in desc
            )
        )

        if is_new_need:
            ui_type = "신규(채번필요)"
            ui_remark = "채번 필요"
        elif ctype == "Changing":
            if _is_meaningful_part_no(base_pno) and _is_meaningful_part_no(new_pno) and base_pno != new_pno:
                ui_type = "변경(대체/사양)"
                ui_remark = f"Base: {base_pno} → New: {new_pno}"
            else:
                ui_type = "변경(구조/형상)"
                ui_remark = ""
        elif ctype == "NEW":
            ui_type = "신규(품번확정)"
            ui_remark = f"P/N: {new_pno}" if _is_meaningful_part_no(new_pno) else ""

    if concern0:
        ui_type = f"{ui_type} · 확인필요" if ui_type else "확인필요"
        ui_remark = f"{ui_remark} | 확인필요: {concern0}" if ui_remark else f"확인필요: {concern0}"

    ui_type = _safe_ui_text(ui_type)
    ui_remark = _sanitize_user_remark_text(_safe_ui_text(ui_remark))
    return ui_type, ui_remark


def _infer_product_type_from_model(model_text: str) -> str:
    m = str(model_text or "").strip().upper()
    if m.startswith("W"):
        return "Built-in Oven"
    if m.startswith("M"):
        return "Microwave Oven"
    return "Oven/Microwave"


PPT_CHANGE_EXTRACTION_PROMPT_TMPL = """
당신은 LG전자 오븐/전자레인지 개발 프로세스 전문가입니다.
사용자가 업로드한 "개발 유형 및 등급 확정 심의회" PPT에서
변경부품리스트 생성에 필요한 정보를 자동 추출합니다.

═══════════════════════════════════════════
[STEP 1] 개발 PJT 메타정보 추출
═══════════════════════════════════════════
PPT 내 "개발 PJT 개요" 영역에서 아래 필드를 추출하세요:

    - 제품군 (예: Built-in Oven)
    - PJT명 (예: 24인치오븐_Extra grade)
    - Base 모델 (예: WSED7667M)
    - 출시 국가 (예: 유럽)
    - 정격 (예: 220V~240V/50Hz)
    - 용량 (예: 76L)
    - 확정 등급 (예: B)
    - 개발 유형 (예: NPI3.0)

출력 형식:
{{
    "project_meta": {{
        "product_type": "...",
        "project_name": "...",
        "base_model": "...",
        "target_country": "...",
        "rating": "...",
        "capacity": "...",
        "dev_grade": "...",
        "dev_type": "..."
    }}
}}

═══════════════════════════════════════════
[STEP 2] 모듈별 변경점 요약 추출
═══════════════════════════════════════════
"유첨1. 개발 변경점 상세 작성 내용" 중
모듈 요약 테이블(No / Module명 / 주요 변경점)에서 추출하세요.

이 테이블은 보통 Exploded View 이미지와 함께 있으며,
①②③... 번호로 모듈이 구분됩니다.

출력 형식:
{{
    "module_summary": [
        {{
            "no": "①",
            "module": "Cavity",
            "changes": [
                "BLDC Conv. Fan Motor 적용 (AC → BLDC) - 30인치 SKS 공용",
                "Fan Blade 신규 적용 (30인치 SKS 공용)",
                "Fan Cover 금형 신작"
            ]
        }}
    ]
}}

═══════════════════════════════════════════
[STEP 3] 상세 변경점 테이블 추출
═══════════════════════════════════════════
"유첨1" 중 아래 컬럼 구조의 상세 테이블에서 추출하세요:
    구분 | No | Part | 변경 내역(변경 전→변경 후) | 변경 사유 | 걱정점

구분은 다음 카테고리로 나뉩니다:
    - 기구 (4M 변경포함)
    - 제어
    - ThinQ App

출력 형식:
{{
    "detailed_changes": [
        {{
            "category": "기구",
            "no": 1,
            "part": "Conv. Motor",
            "change_detail": "AC motor → BLDC motor",
            "change_reason": "균일 가열 성능 및 요리 시간 단축",
            "concern": "요리 성능 및 온도 정밀도",
            "tags": ["BLDC", "모터"]
        }}
    ]
}}

tags 필드 규칙:
- 변경 내역과 Part명에서 핵심 키워드를 자동 태깅하세요.
- 태그는 나중에 BOM 검색 시 매칭 키워드로 사용됩니다.
- 빈 행(Part, 변경 내역이 모두 비어있는 행)은 무시하세요.

═══════════════════════════════════════════
[STEP 4] 모듈별 상세 변경점 추출 (부품번호 포함)
═══════════════════════════════════════════
"유첨1. 개발 변경점_XXX Assembly" 슬라이드들에서 추출하세요.
이 슬라이드들은 Module | Base | New model | Remark 구조입니다.

출력 형식:
{{
    "module_details": [
        {{
            "module": "Cavity",
            "sub_module": "Convection Fan Motor",
            "base": {{
                "description": "AC Motor + Bracket → CCW",
                "part_no": "EAU65078501"
            }},
            "new": {{
                "description": "BLDC Motor + Bracket (30인치 SKS 공용) → 정/역회전, 가변풍량",
                "part_no": "4810W1N060B",
                "is_shared": true,
                "shared_with": "30인치 SKS"
            }},
            "remark": "Inverter Pro Bake BLDC 모터 개발 - CMR 요리성능 개선"
        }}
    ]
}}

is_shared 판별 규칙:
- "공용", "공용화", "수평전개", "동일 적용" 등의 키워드가 있으면
    is_shared = true, shared_with에 공용 대상 모델/사이즈를 기록
- 해당 키워드가 없으면 is_shared = false

═══════════════════════════════════════════
[STEP 5] 예상 우려점 추출
═══════════════════════════════════════════
"유첨2. 개발 예상 우려점 검토 결과" 테이블에서 추출하세요.
구분: 기구/구조 부품 변경점 | 전장/PCB품 변경점 | 부자재 등 기타 변경점

출력 형식:
{{
    "risk_review": [
        {{
            "category": "기구/구조 부품 변경점",
            "concern": "...",
            "countermeasure": "...",
            "attachment": "..."
        }}
    ]
}}
※ 내용이 비어있는 경우 "내용 없음"으로 표기하세요.

═══════════════════════════════════════════
[STEP 6] 변경부품리스트 생성용 변경점 정리
═══════════════════════════════════════════
위 STEP 2~4의 결과를 종합하여,
변경부품리스트 자동 생성 시스템에 전달할 최종 변경점 목록을
아래 형식으로 정리하세요:

{{
    "change_points_for_bom": [
        {{
            "id": 1,
            "description": "Convection Motor BLDC 적용 (AC → BLDC)",
            "module": "Cavity",
            "type": "스펙변경",
            "base_part_no": "EAU65078501",
            "new_part_no": "4810W1N060B",
            "is_shared_part": true,
            "shared_source": "30인치 SKS",
            "related_parts": [
                "Fan Blade (MDG63965901)",
                "Fan Cover (MCK71660202 → 신규)"
            ],
            "concerns": ["요리 성능 및 온도 정밀도"]
        }}
    ]
}}

type 분류 기준:
- 스펙변경: 기존 부품의 사양이 변경 (예: AC→BLDC, 4.3"→6.8")
- 신규추가: 기존에 없던 부품/모듈 추가 (예: Camera 추가)
- 구조변경: 조립 구조, 체결 방식 변경 (예: Bracket 개조)
- 사이즈변경: 치수/크기 변경

is_shared_part가 true인 부품은
변경부품리스트에서 "공용 적용"으로 표시하고,
신규 개발이 아님을 명시하세요.

═══════════════════════════════════════════
[주의사항]
═══════════════════════════════════════════
1. PPT에서 텍스트를 추출할 때 취소선(~~텍스트~~)이 있으면
     해당 내용은 삭제된 것으로 간주하고 무시하세요.
     취소선 아래 새로 작성된 내용을 최종값으로 사용
2. "변경 없음" 태그가 있는 슬라이드는 해당 모듈의
     기존 내용이 유지됨을 의미합니다.
3. "추가 변경점" 태그가 있는 슬라이드는
     기존 심의 후 추가된 변경점이므로 반드시 포함하세요.
4. 이미지, Feature List 등 비텍스트 첨부는
     존재 여부만 기록하세요 (내용 추출 불필요).
5. 추출한 정보에 대해 절대 임의로 부품번호나 스펙을 생성하지 마세요.
     PPT에 없는 정보는 "TBD" 또는 "정보 없음"으로 표기하세요.

아래 프로젝트 컨텍스트를 참고하세요:
- 제품군(추정): {product_type}
- Base 모델(세션): {source_model}
- Target 모델(세션): {target_model}
- 주요 차이(세션): {key_diff}
- 개발등급(세션): {dev_grade}
- 변경점(세션):
{change_points}
""".strip()


def build_ppt_change_extraction_prompt(
    change_items: list[str],
    source_model: str,
    target_model: str,
    key_diff: str,
    dev_grade: str,
    product_type: str = "",
) -> str:
    ptype = str(product_type or "").strip() or _infer_product_type_from_model(target_model)
    points = [f"- {str(x).strip()}" for x in (change_items or []) if str(x or "").strip()]
    if not points:
        points = ["- (변경점 없음)"]
    return PPT_CHANGE_EXTRACTION_PROMPT_TMPL.format(
        product_type=ptype,
        source_model=str(source_model or ""),
        target_model=str(target_model or ""),
        key_diff=str(key_diff or ""),
        dev_grade=str(dev_grade or ""),
        change_points="\n".join(points),
    )


def _split_lines_text(v: str) -> list[str]:
    txt = str(v or "")
    parts = re.split(r"[\n\r•·]+", txt)
    out = []
    for p in parts:
        s = re.sub(r"\s+", " ", p).strip(" -\t")
        if s:
            out.append(s)
    return out


def _extract_part_no_candidates(text: str) -> list[str]:
    cands = re.findall(r"\b[A-Z0-9]{8,14}\b", str(text or "").upper())
    out = []
    for c in cands:
        if c not in out:
            out.append(c)
    return out


def _strip_part_no_from_desc(text: str, pnos: list[str]) -> str:
    s = str(text or "")
    for p in pnos or []:
        s = re.sub(re.escape(p), "", s, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", s).strip(" -:|\t")


def _guess_change_type(part: str, detail: str, base_desc: str, new_desc: str, remark: str) -> str:
    t = " ".join([part, detail, base_desc, new_desc, remark]).lower()
    if any(k in t for k in ["inch", "인치", "사이즈", "size", "mm", "치수", "폭", "높이"]):
        return "사이즈변경"
    if any(k in t for k in ["bracket", "체결", "구조", "layout", "개조", "frame", "조립"]):
        return "구조변경"
    if any(k in t for k in ["신규", "추가", "add", "장착", "new"]):
        return "신규추가"
    return "스펙변경"


def _auto_tags(part: str, detail: str, limit: int = 6) -> list[str]:
    txt = f"{part} {detail}".upper()
    found = []
    seed = [
        "BLDC", "AC", "DC", "MOTOR", "FAN", "CAMERA", "HARNESS", "PCB", "UI",
        "BRACKET", "INVERTER", "HEATER", "SENSOR", "DISPLAY", "TOUCH", "THINQ",
        "도어", "카메라", "모터", "하네스", "제어", "앱", "팬", "브라켓", "센서",
    ]
    for s in seed:
        if s in txt and s not in found:
            found.append(s)
    toks = re.findall(r"[A-Z][A-Z0-9\-]{1,15}|[가-힣]{2,12}", txt)
    stop = {"PART", "CHANGE", "MODEL", "BASE", "NEW", "NO", "AND", "THE", "변경", "내역"}
    for t in toks:
        if t in stop or t.isdigit() or t in found:
            continue
        found.append(t)
        if len(found) >= limit:
            break
    return found[:limit]


def _ppt_text_from_shape(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    chunks = []
    tf = shape.text_frame
    for p in tf.paragraphs:
        run_text = []
        for r in p.runs:
            try:
                if getattr(r.font, "strike", False):
                    continue
            except Exception:
                pass
            run_text.append(r.text or "")
        line = "".join(run_text).strip()
        if line:
            chunks.append(line)
    return "\n".join(chunks).strip()


def _norm_header(h: str) -> str:
    return re.sub(r"[^A-Z0-9가-힣]", "", str(h or "").upper())


def _sanitize_product_type(v: Any) -> str:
    s = re.sub(r"\s+", " ", str(v or "")).strip(" -:\t")
    if not s:
        return ""
    bad_markers = ["동일 사업부", "다르면", "으로 봄", "EX:", "예:", "정의", "설명", "기준", "참고", "해당"]
    su = s.upper()
    if any(b.upper() in su for b in bad_markers):
        return ""
    if any(x in s for x in [".", "!", "?", ";", "(ex", "예)"]):
        return ""
    if " " in s and len(s) > 24:
        return ""
    if len(s) > 40:
        return ""
    return s


def _extract_product_type_from_page1(text_p1: str) -> str:
    lines = [re.sub(r"\s+", " ", x).strip() for x in str(text_p1 or "").splitlines() if str(x or "").strip()]
    cands: list[str] = []

    # 1) key:value 또는 key|value 패턴 우선
    p1 = [
        r"(?:^|\|)\s*제품군\s*(?:\||[:：])\s*([^|\n\r]+)",
        r"(?:^|\|)\s*PRODUCT(?:\s*TYPE)?\s*(?:\||[:：])\s*([^|\n\r]+)",
    ]
    for p in p1:
        for m in re.finditer(p, text_p1, flags=re.IGNORECASE):
            v = _sanitize_product_type(m.group(1))
            if v:
                cands.append(v)

    # 2) 테이블 행 형태: ... | 제품군 | 값 | ...
    for ln in lines:
        if "제품군" in ln or "PRODUCT" in ln.upper():
            toks = [t.strip() for t in ln.split("|") if t.strip()]
            for i, t in enumerate(toks):
                tu = t.upper()
                if t == "제품군" or tu in ("PRODUCT", "PRODUCT TYPE"):
                    if i + 1 < len(toks):
                        v = _sanitize_product_type(toks[i + 1])
                        if v:
                            cands.append(v)

    # 2-1) 인접 라인 패턴: "제품군" 다음 줄이 값인 경우
    for i, ln in enumerate(lines):
        lnu = ln.upper()
        is_key_only = _norm_header(ln) in {"제품군", "PRODUCT", "PRODUCTTYPE"}
        if is_key_only or ("제품군" in ln) or ("PRODUCT TYPE" in lnu):
            if i + 1 < len(lines):
                nxt = lines[i + 1]
                # 메타 키로 보이는 라인은 건너뜀
                if _norm_header(nxt) not in {"PJT명", "PROJECTNAME", "BASE모델", "BASEMODEL", "개발등급", "확정등급", "개발유형"}:
                    v = _sanitize_product_type(nxt)
                    if v:
                        cands.append(v)

    # 3) 후보 점수화: 짧고 제품명 키워드가 있는 값 우선
    def _score(v: str) -> int:
        s = str(v or "").strip()
        su = s.upper()
        score = 0
        # BUILT-IN을 최우선으로 처리
        if "BUILT-IN" in su:
            score += 50
        if any(k in su for k in ["OVEN", "MICROWAVE", "레인지", "전자레인지", "오븐"]):
            score += 10
        if len(s) <= 20:
            score += 4
        if " " not in s:
            score += 2
        if any(k in su for k in ["동일사업부", "다르면", "설명", "정의", "EX", "예"]):
            score -= 100
        return score

    cands = [c for c in cands if c]
    if not cands:
        return ""
    cands = sorted(cands, key=lambda x: (_score(x), -len(x)), reverse=True)
    return cands[0]


def _extract_product_type_from_anywhere(text: str) -> str:
    lines = [re.sub(r"\s+", " ", x).strip() for x in str(text or "").splitlines() if str(x or "").strip()]
    cands: list[str] = []

    # 제품군/PRODUCT TYPE 키 기준 키|값 형태만 허용 (자유문장 하드코딩 금지)
    for ln in lines:
        up = ln.upper()
        if ("제품군" not in ln) and ("PRODUCT" not in up):
            continue
        toks = [t.strip() for t in ln.split("|") if t.strip()]
        for i, t in enumerate(toks):
            tu = t.upper()
            if t == "제품군" or tu in ("PRODUCT", "PRODUCT TYPE"):
                if i + 1 < len(toks):
                    v = _sanitize_product_type(toks[i + 1])
                    if v:
                        cands.append(v)

    # 인접 라인 패턴 허용: KEY 라인 다음 라인 값
    for i, ln in enumerate(lines):
        lnu = ln.upper()
        is_key_only = _norm_header(ln) in {"제품군", "PRODUCT", "PRODUCTTYPE"}
        if is_key_only or ("제품군" in ln) or ("PRODUCT TYPE" in lnu):
            if i + 1 < len(lines):
                v = _sanitize_product_type(lines[i + 1])
                if v:
                    cands.append(v)

    def _score(v: str) -> int:
        s = str(v or "").strip()
        su = str(v or "").upper()
        score = 0
        if any(k in su for k in ["OVEN", "MICROWAVE", "BUILT-IN", "BUILT IN", "RANGE", "COOKTOP", "REFRIGERATOR", "WASHER", "DRYER", "DISHWASHER"]):
            score += 10
        if len(s) <= 24:
            score += 3
        if any(k in su for k in ["동일사업부", "다르면", "설명", "정의", "EX", "예", "기준", "참고"]):
            score -= 100
        return score

    cands = [c for c in cands if c]
    if not cands:
        return ""
    cands.sort(key=lambda x: (_score(x), -len(x)), reverse=True)
    return cands[0]


def _classify_slide_role(slide_lines: list[str], table_headers: list[str]) -> dict:
    text = "\n".join(slide_lines or [])
    hdr = " | ".join(table_headers or [])
    blob = f"{text}\n{hdr}".upper()

    overview_hits = 0
    detail_hits = 0
    ref_hits = 0

    if "개발 PJT 개요".upper() in blob:
        overview_hits += 3
    if "심의 결과".upper() in blob or "확정".upper() in blob:
        overview_hits += 3
    for k in ["개발 유형", "개발 등급", "제품군", "BASE MODEL", "BASE 모델", "PJT명", "PROJECT NAME"]:
        if k.upper() in blob:
            overview_hits += 1

    for k in ["유첨", "개발 변경점 상세", "CAVITY", "DOOR", "CONTROLLER ASSEMBLY", "구분", "변경 내역", "BASE", "NEW"]:
        if k.upper() in blob:
            detail_hits += 1

    for k in ["개발 등급 분류 기준", "운영 기준", "부표", "기준", "절차", "정의", "상세", "참고"]:
        if k.upper() in blob:
            ref_hits += 1

    # 최우선 규칙: 개요 키워드 동시 출현은 신뢰도 최고
    if "개발 PJT 개요".upper() in blob and "심의 결과".upper() in blob:
        return {"role": "OVERVIEW", "reason": "overview_high_conf"}

    if overview_hits >= max(4, detail_hits + 1) and overview_hits >= ref_hits:
        return {"role": "OVERVIEW", "reason": "overview_score"}

    if detail_hits >= max(2, ref_hits + 1):
        return {"role": "DETAIL", "reason": "detail_score"}

    if ref_hits >= 2:
        return {"role": "REFERENCE", "reason": "reference_score"}

    return {"role": "UNKNOWN", "reason": "weak_signal"}


def _extract_project_meta(lines: list[str], ctx: dict, page1_lines: list[str] | None = None) -> dict:
    text = "\n".join(lines)
    text_p1 = "\n".join(page1_lines or [])
    text_for_ptype = f"{text_p1}\n{text}"
    row = {
        "product_type": "정보 없음",
        "project_name": "정보 없음",
        "base_model": str(ctx.get("source_model") or "정보 없음"),
        "target_country": "정보 없음",
        "rating": "정보 없음",
        "capacity": "정보 없음",
        "dev_grade": str(ctx.get("dev_grade") or "정보 없음"),
        "dev_type": "정보 없음",
    }
    kv_patterns = {
        "project_name": [r"PJT명\s*[:：]\s*(.+)", r"PROJECT\s*NAME\s*[:：]\s*(.+)"],
        "base_model": [r"BASE\s*모델\s*[:：]\s*([A-Z0-9]{8,14})", r"BASE\s*MODEL\s*[:：]\s*([A-Z0-9]{8,14})"],
        "target_country": [r"출시\s*국가\s*[:：]\s*(.+)", r"COUNTRY\s*[:：]\s*(.+)"],
        "rating": [r"정격\s*[:：]\s*(.+)", r"RATING\s*[:：]\s*(.+)"],
        "capacity": [r"용량\s*[:：]\s*(.+)", r"CAPACITY\s*[:：]\s*(.+)"],
        "dev_grade": [r"확정\s*등급\s*[:：]\s*([A-Z0-9\.\-]+)", r"등급\s*[:：]\s*([A-Z0-9\.\-]+)"],
        "dev_type": [r"개발\s*유형\s*[:：]\s*([A-Z0-9\.\-]+)", r"DEV\s*TYPE\s*[:：]\s*([A-Z0-9\.\-]+)"],
    }

    def _clean_meta_value(v: Any) -> str:
        s = re.sub(r"\s+", " ", str(v or "")).strip(" -:\t")
        if not s:
            return ""
        if s.upper() in {"TBD", "N/A", "NA", "-", "X", "O"}:
            return ""
        if any(k in s for k in ["정보 없음", "내용 없음"]):
            return ""
        return s

    def _extract_from_pipe_lines(src_lines: list[str], key_aliases: list[str]) -> str:
        alias_norm = [_norm_header(a) for a in key_aliases]
        for ln in src_lines or []:
            toks = [re.sub(r"\s+", " ", t).strip() for t in str(ln or "").split("|") if str(t or "").strip()]
            if len(toks) < 2:
                continue
            for i, t in enumerate(toks):
                tn = _norm_header(t)
                if any(a and (a == tn or a in tn) for a in alias_norm):
                    if i + 1 < len(toks):
                        cand = _clean_meta_value(toks[i + 1])
                        if cand:
                            return cand
        return ""

    for k, pats in kv_patterns.items():
        for p in pats:
            m = re.search(p, text, flags=re.IGNORECASE)
            if m and str(m.group(1)).strip():
                row[k] = str(m.group(1)).strip()
                break

    # 표 형태(키|값)에서 메타를 추출한다. page1을 최우선으로 보고 없으면 전체 텍스트를 본다.
    p1_lines = [x for x in str(text_p1 or "").splitlines() if str(x or "").strip()]
    all_src_lines = [x for x in str(text or "").splitlines() if str(x or "").strip()]
    table_keys = {
        "project_name": ["PJT명", "PROJECT NAME", "프로젝트명"],
        "base_model": ["BASE 모델", "BASE MODEL", "BASE"],
        "target_country": ["출시 국가", "COUNTRY", "적용 국가"],
        "rating": ["정격", "RATING"],
        "capacity": ["용량", "CAPACITY"],
        "dev_grade": ["확정 등급", "개발 등급", "등급"],
        "dev_type": ["개발 유형", "DEV TYPE", "유형"],
    }
    for k, aliases in table_keys.items():
        if row.get(k) not in ("", "정보 없음"):
            continue
        cand = _extract_from_pipe_lines(p1_lines, aliases) or _extract_from_pipe_lines(all_src_lines, aliases)
        if cand:
            row[k] = cand

    # product_type는 1페이지 템플릿만 source-of-truth로 사용한다.
    p1_ptype = _extract_product_type_from_page1(text_p1)
    if p1_ptype:
        row["product_type"] = p1_ptype

    # 1페이지에서 못 잡으면 전체 텍스트에서 Built-in Oven/제품군 키를 보강 추출
    if row["product_type"] in ("", "정보 없음"):
        any_ptype = _extract_product_type_from_anywhere(text_for_ptype)
        if any_ptype:
            row["product_type"] = any_ptype

    if row["product_type"] == "정보 없음":
        # target_model이 비어있는 경우가 많아서 source_model/base_model까지 순차 fallback
        model_hint = str(
            ctx.get("target_model")
            or ctx.get("source_model")
            or row.get("base_model")
            or ""
        )
        row["product_type"] = str(_sanitize_product_type(ctx.get("product_type")) or _infer_product_type_from_model(model_hint))
    return row


def extract_change_review_from_pptx_bytes(pptx_bytes: bytes, ctx: dict | None = None) -> dict:
    ctx = ctx or {}
    try:
        from pptx import Presentation
    except Exception as e:
        raise RuntimeError("python-pptx 패키지가 필요합니다. 설치 후 다시 시도하세요.") from e

    import io
    prs = Presentation(io.BytesIO(pptx_bytes))

    discarded: list[dict] = []
    all_lines: list[str] = []
    all_table_text_lines: list[str] = []
    slide_lines_map: dict[int, list[str]] = {}
    slide_table_headers_map: dict[int, list[str]] = {}
    slide_table_text_map: dict[int, list[str]] = {}
    table_rows: list[dict] = []
    image_presence: list[dict] = []

    def _flat_cells(rows: list[list[str]]) -> list[str]:
        out = []
        for r in rows or []:
            for c in r or []:
                t = re.sub(r"\s+", " ", str(c or "")).strip()
                if t:
                    out.append(t)
        return out

    def _classify_table_candidate(rows: list[list[str]], slide_lines: list[str], slide: int, table_id: int) -> dict:
        if not rows:
            return {"type": "IGNORE", "why": "empty_table"}

        raw_header = [re.sub(r"\s+", " ", str(x or "")).strip() for x in (rows[0] or [])]
        header_norm = [_norm_header(x) for x in raw_header]
        cells = _flat_cells(rows)
        all_text = " ".join(raw_header + cells + (slide_lines or []))

        # 블랙리스트 키워드: 테이블 성격 자체 제외
        blk = ["적용여부", "과전압", "PDR", "ODR", "EVENT", "NPI", "ACTIVITY"]
        all_up = all_text.upper()
        if any(k.upper() in all_up for k in blk):
            return {"type": "IGNORE", "why": "blacklist_keyword"}

        # O/X 체크표 비율이 높은 경우 제외
        chk_tokens = {"O", "X", "OX", "(O,X)", "(O,X)", "(O, X)"}
        non_empty = [c for c in cells if str(c).strip()]
        chk_cnt = 0
        for c in non_empty:
            n = re.sub(r"\s+", "", str(c).upper())
            if n in chk_tokens:
                chk_cnt += 1
        if non_empty and (chk_cnt / max(len(non_empty), 1)) >= 0.20 and len(non_empty) >= 4:
            return {"type": "IGNORE", "why": "checkbox_ratio_high"}

        hraw_join = " | ".join(raw_header).upper()
        hnorm_join = "|".join(header_norm)
        hraw_compact = re.sub(r"[^A-Z0-9가-힣]", "", hraw_join.upper())
        hnorm_compact = re.sub(r"[^A-Z0-9가-힣]", "", hnorm_join.upper())

        def _has_any(src: str, kws: list[str]) -> bool:
            s = src.upper()
            return any(k.upper() in s for k in kws)

        def _has_any_compact(src_compact: str, kws: list[str]) -> bool:
            return any(re.sub(r"[^A-Z0-9가-힣]", "", k.upper()) in src_compact for k in kws)

        # MAIN_CHANGE_TABLE
        has_change = _has_any(hraw_join, ["변경내역", "변경 내역", "CHANGE"]) or _has_any_compact(hraw_compact, ["변경내역"])
        has_before_after = (
            _has_any(hraw_join, ["변경 전", "변경 후", "->", "→", "BEFORE", "AFTER"])
            or _has_any_compact(hraw_compact, ["변경전", "변경후"])
            or ("변경전" in hnorm_join or "변경후" in hnorm_join)
        )
        has_reason = _has_any(hraw_join, ["변경사유", "변경 사유", "REASON"]) or _has_any_compact(hraw_compact, ["변경사유"])
        has_concern = _has_any(hraw_join, ["걱정점", "걱정 점", "우려", "CONCERN"]) or _has_any_compact(hraw_compact, ["걱정점", "우려"])
        has_part = _has_any(hraw_join, ["PART", "부품"]) or _has_any_compact(hraw_compact, ["PART", "부품"])
        if has_change and has_before_after and has_reason and has_concern and has_part:
            return {"type": "MAIN_CHANGE_TABLE", "why": "whitelist_main_change"}

        # DETAIL_CHANGE_TABLE
        has_base = _has_any(hraw_join, ["BASE"])
        has_new = _has_any(hraw_join, ["NEW"])
        has_mod_or_remark = _has_any(hraw_join, ["MODULE", "모듈", "REMARK", "비고"])
        if has_base and has_new and has_mod_or_remark:
            return {"type": "DETAIL_CHANGE_TABLE", "why": "whitelist_detail_change"}

        # RISK_TABLE
        has_risk = _has_any(hraw_join, ["우려", "걱정", "CONCERN"]) or _has_any_compact(hraw_compact, ["우려", "걱정"])
        has_counter = _has_any(hraw_join, ["대책", "대응", "검증", "방안", "COUNTER"]) or _has_any_compact(hraw_compact, ["대책", "대응", "검증방안"])
        if has_risk and has_counter:
            return {"type": "RISK_TABLE", "why": "whitelist_risk"}

        return {"type": "IGNORE", "why": "not_whitelisted"}

    for s_idx, slide in enumerate(prs.slides, 1):
        slide_lines = []
        slide_table_headers: list[str] = []
        slide_table_texts: list[str] = []
        has_non_text_attach = False
        table_seq = 0
        for sh in slide.shapes:
            if getattr(sh, "has_text_frame", False):
                txt = _ppt_text_from_shape(sh)
                if txt:
                    ls = _split_lines_text(txt)
                    slide_lines.extend(ls)
                    all_lines.extend(ls)
            if getattr(sh, "shape_type", None) in (13, 14):
                has_non_text_attach = True
            if getattr(sh, "has_table", False):
                table_seq += 1
                t = sh.table
                rows = []
                for r in t.rows:
                    row = []
                    for c in r.cells:
                        row.append(re.sub(r"\s+", " ", c.text or "").strip())
                    rows.append(row)
                if rows:
                    for rr in rows:
                        rrj = " | ".join([re.sub(r"\s+", " ", str(x or "")).strip() for x in rr if str(x or "").strip()])
                        if rrj:
                            slide_table_texts.append(rrj)
                            all_table_text_lines.append(rrj)
                    hdr_join = " | ".join([re.sub(r"\s+", " ", str(x or "")).strip() for x in (rows[0] or []) if str(x or "").strip()])
                    if hdr_join:
                        slide_table_headers.append(hdr_join)
                    table_rows.append({
                        "slide": s_idx,
                        "table_id": f"S{s_idx}-T{table_seq}",
                        "rows": rows,
                        "slide_lines": list(slide_lines),
                    })
        slide_lines_map[s_idx] = list(slide_lines)
        slide_table_headers_map[s_idx] = list(slide_table_headers)
        slide_table_text_map[s_idx] = list(slide_table_texts)
        image_presence.append({"slide": s_idx, "has_attachment": bool(has_non_text_attach)})

    slide_roles: dict[int, str] = {}
    slide_role_details: list[dict] = []
    for s_idx in sorted(slide_lines_map.keys()):
        cls = _classify_slide_role(slide_lines_map.get(s_idx) or [], slide_table_headers_map.get(s_idx) or [])
        role = str(cls.get("role") or "UNKNOWN")
        slide_roles[s_idx] = role
        slide_role_details.append({"slide": s_idx, "role": role, "reason": cls.get("reason", "")})

    overview_scope_lines: list[str] = []
    for s_idx, role in slide_roles.items():
        if role == "OVERVIEW":
            overview_scope_lines.extend(slide_lines_map.get(s_idx) or [])
            overview_scope_lines.extend(slide_table_text_map.get(s_idx) or [])
            overview_scope_lines.extend(slide_table_headers_map.get(s_idx) or [])

    # 개요/결론 슬라이드가 없으면 약한 fallback: 첫 슬라이드 사용
    if not overview_scope_lines and slide_lines_map:
        first_idx = sorted(slide_lines_map.keys())[0]
        overview_scope_lines.extend(slide_lines_map.get(first_idx) or [])
        overview_scope_lines.extend(slide_table_text_map.get(first_idx) or [])
        overview_scope_lines.extend(slide_table_headers_map.get(first_idx) or [])

    # 제품군/프로젝트 메타는 첫 슬라이드 템플릿을 source-of-truth로 우선 사용한다.
    page1_scope_lines: list[str] = []
    if slide_lines_map:
        first_idx = sorted(slide_lines_map.keys())[0]
        page1_scope_lines.extend(slide_lines_map.get(first_idx) or [])
        page1_scope_lines.extend(slide_table_text_map.get(first_idx) or [])
        page1_scope_lines.extend(slide_table_headers_map.get(first_idx) or [])

    # project_meta는 OVERVIEW + table cell 텍스트 중심으로 읽고, 전체 텍스트는 보조로 합친다.
    meta_scope_lines: list[str] = []
    meta_scope_lines.extend(overview_scope_lines)
    meta_scope_lines.extend(all_table_text_lines)
    meta_scope_lines.extend(all_lines)

    project_meta = _extract_project_meta(meta_scope_lines, ctx, page1_lines=page1_scope_lines or overview_scope_lines)

    module_summary: list[dict] = []
    detailed_changes: list[dict] = []
    module_details: list[dict] = []
    risk_review: list[dict] = []
    module_summary_hints: list[dict] = []

    for tb in table_rows:
        rows = tb["rows"]
        slide_no = int(tb.get("slide") or 0)
        slide_role = slide_roles.get(slide_no, "UNKNOWN")

        # 부표/기준 슬라이드는 결정 정보에서 제외
        if slide_role == "REFERENCE":
            flat = _flat_cells(rows)
            discarded.append({
                "why": "reference_slide_scope",
                "slide": tb.get("slide"),
                "table_id": tb.get("table_id"),
                "sample_text": " | ".join(flat[:6]) if flat else "",
            })
            continue

        cls = _classify_table_candidate(rows, tb.get("slide_lines") or [], tb.get("slide"), tb.get("table_id"))
        ttype = cls.get("type", "IGNORE")

        if ttype == "IGNORE":
            # DETAIL 슬라이드의 "Module명 | 주요 변경점" 요약표는 module 매핑 힌트로만 사용
            if slide_role == "DETAIL" and rows and rows[0]:
                hdr0 = [_norm_header(x) for x in rows[0]]
                has_mod_col = any(("MODULE" in h) or ("모듈" in h) for h in hdr0)
                has_change_col = any(("주요변경점" in h) or ("변경내역" in h) or ("CHANGE" in h) for h in hdr0)
                if has_mod_col and has_change_col:
                    m_i = next((i for i, h in enumerate(hdr0) if ("MODULE" in h or "모듈" in h)), 0)
                    c_i = next((i for i, h in enumerate(hdr0) if ("주요변경점" in h or "변경내역" in h or "CHANGE" in h)), min(1, len(hdr0)-1))
                    cur_mod = ""
                    for rr in rows[1:]:
                        mod = (rr[m_i] if m_i < len(rr) else "").strip() or cur_mod
                        cur_mod = mod or cur_mod
                        chg = (rr[c_i] if c_i < len(rr) else "").strip()
                        if mod and chg:
                            module_summary_hints.append({
                                "module": mod,
                                "text": chg,
                                "src": {"slide": tb.get("slide"), "table_id": tb.get("table_id")},
                            })

            flat = _flat_cells(rows)
            discarded.append({
                "why": cls.get("why", "ignored"),
                "slide": tb.get("slide"),
                "table_id": tb.get("table_id"),
                "sample_text": " | ".join(flat[:6]) if flat else "",
            })
            continue

        hdr = [_norm_header(x) for x in rows[0]]

        # 상세 변경 추출은 상세 슬라이드에서만 허용
        if slide_role != "DETAIL" and ttype in ("MAIN_CHANGE_TABLE", "DETAIL_CHANGE_TABLE", "RISK_TABLE"):
            flat = _flat_cells(rows)
            discarded.append({
                "why": "non_detail_slide_for_change_table",
                "slide": tb.get("slide"),
                "table_id": tb.get("table_id"),
                "sample_text": " | ".join(flat[:6]) if flat else "",
            })
            continue

        # MAIN_CHANGE_TABLE => STEP3 only
        if ttype == "MAIN_CHANGE_TABLE":
            c_i = next((i for i, h in enumerate(hdr) if ("구분" in h or "CATEGORY" in h)), 0)
            n_i = next((i for i, h in enumerate(hdr) if "NO" in h), 1)
            p_i = next((i for i, h in enumerate(hdr) if ("PART" in h or "부품" in h)), 2)
            d_i = next((i for i, h in enumerate(hdr) if ("변경내역" in h or "CHANGE" in h)), 3)
            r_i = next((i for i, h in enumerate(hdr) if ("변경사유" in h or "REASON" in h)), 4)
            g_i = next((i for i, h in enumerate(hdr) if ("걱정" in h or "CONCERN" in h)), 5)
            cur_cat = ""
            for r in rows[1:]:
                cat = (r[c_i] if c_i < len(r) else "").strip() or cur_cat
                cur_cat = cat or cur_cat
                no = (r[n_i] if n_i < len(r) else "").strip()
                part = (r[p_i] if p_i < len(r) else "").strip()
                detail = (r[d_i] if d_i < len(r) else "").strip()
                reason = (r[r_i] if r_i < len(r) else "").strip()
                concern = (r[g_i] if g_i < len(r) else "").strip()
                if not part and not detail:
                    discarded.append({
                        "why": "empty_part_and_detail_row",
                        "slide": tb.get("slide"),
                        "table_id": tb.get("table_id"),
                        "sample_text": " | ".join([str(x) for x in r[:6]]),
                    })
                    continue
                n_val: Any = no
                try:
                    n_val = int(re.sub(r"[^0-9]", "", no)) if re.search(r"[0-9]", no) else no
                except Exception:
                    n_val = no
                detailed_changes.append({
                    "category": cat or "정보 없음",
                    "discipline": cat or "정보 없음",
                    "no": n_val,
                    "part": part or "정보 없음",
                    "change_detail": detail or "정보 없음",
                    "change_reason": reason or "정보 없음",
                    "concern": concern or "내용 없음",
                    "tags": _auto_tags(part, detail),
                    "_src": {"slide": tb.get("slide"), "table_id": tb.get("table_id")},
                })
            continue

        # DETAIL_CHANGE_TABLE => STEP4 only
        if ttype == "DETAIL_CHANGE_TABLE":
            m_i = next((i for i, h in enumerate(hdr) if ("MODULE" in h or "모듈" in h)), 0)
            b_i = next((i for i, h in enumerate(hdr) if "BASE" in h), 1)
            n_i = next((i for i, h in enumerate(hdr) if "NEW" in h), 2)
            rm_i = next((i for i, h in enumerate(hdr) if ("REMARK" in h or "비고" in h)), min(3, len(hdr)-1))
            for r in rows[1:]:
                mod = (r[m_i] if m_i < len(r) else "").strip()
                base_txt = (r[b_i] if b_i < len(r) else "").strip()
                new_txt = (r[n_i] if n_i < len(r) else "").strip()
                remark = (r[rm_i] if rm_i < len(r) else "").strip()
                if not mod and not base_txt and not new_txt:
                    continue

                # 상세표 행 안전망: 체크표류 차단
                mod_up = mod.upper()
                chk_desc = re.sub(r"\s+", "", f"{new_txt or base_txt}").upper()
                if "과전압" in mod_up or chk_desc in {"O", "X", "(O,X)", "OX"}:
                    discarded.append({
                        "why": "detail_row_blacklist_or_checkbox",
                        "slide": tb.get("slide"),
                        "table_id": tb.get("table_id"),
                        "sample_text": f"{mod} | {base_txt} | {new_txt}",
                    })
                    continue

                base_pnos = _extract_part_no_candidates(base_txt)
                new_pnos = _extract_part_no_candidates(new_txt)
                mix_txt = f"{new_txt} {remark}".lower()
                is_shared = any(k in mix_txt for k in ["공용", "공용화", "수평전개", "동일 적용"])
                sw = ""
                if is_shared:
                    m_sw = re.search(r"(\d{2}\s*인치\s*[A-Z0-9가-힣 ]+|[A-Z0-9가-힣 ]+\s*공용)", f"{new_txt} {remark}", flags=re.IGNORECASE)
                    if m_sw:
                        sw = re.sub(r"\s+", " ", m_sw.group(1)).strip()
                module_details.append({
                    "module": mod or "정보 없음",
                    "sub_module": mod or "정보 없음",
                    "base": {
                        "description": _strip_part_no_from_desc(base_txt, base_pnos) or (base_txt or "정보 없음"),
                        "part_no": base_pnos[0] if base_pnos else "TBD",
                    },
                    "new": {
                        "description": _strip_part_no_from_desc(new_txt, new_pnos) or (new_txt or "정보 없음"),
                        "part_no": new_pnos[0] if new_pnos else "TBD",
                        "is_shared": bool(is_shared),
                        "shared_with": sw,
                    },
                    "remark": remark or "정보 없음",
                    "_src": {"slide": tb.get("slide"), "table_id": tb.get("table_id")},
                })
            continue

        # RISK_TABLE => STEP5 only
        if ttype == "RISK_TABLE":
            c_i = next((i for i, h in enumerate(hdr) if ("구분" in h or "CATEGORY" in h)), 0)
            o_i = next((i for i, h in enumerate(hdr) if ("우려" in h or "CONCERN" in h)), 1)
            m_i = next((i for i, h in enumerate(hdr) if ("대책" in h or "COUNTER" in h)), 2)
            a_i = next((i for i, h in enumerate(hdr) if ("첨부" in h or "ATTACH" in h)), min(3, len(hdr)-1))
            for r in rows[1:]:
                cat = (r[c_i] if c_i < len(r) else "").strip() or "정보 없음"
                concern = (r[o_i] if o_i < len(r) else "").strip() or "내용 없음"
                cm = (r[m_i] if m_i < len(r) else "").strip() or "내용 없음"
                at = (r[a_i] if a_i < len(r) else "").strip() or "내용 없음"
                if cat == "정보 없음" and concern == "내용 없음" and cm == "내용 없음":
                    continue
                risk_review.append({
                    "category": cat,
                    "concern": concern,
                    "countermeasure": cm,
                    "attachment": at,
                })

    # STEP2 (요약): 상세 변경점을 모듈/카테고리 기준으로 합성
    mod_map: dict[str, list[str]] = {}
    for dc in detailed_changes:
        mod = str(dc.get("category") or "정보 없음").strip()
        msg = str(dc.get("change_detail") or "").strip()
        if not msg:
            continue
        mod_map.setdefault(mod, [])
        if msg not in mod_map[mod]:
            mod_map[mod].append(msg)
    module_summary = [
        {"no": str(i + 1), "module": k, "changes": v}
        for i, (k, v) in enumerate(mod_map.items())
    ]

    # STEP6: MAIN_CHANGE_TABLE(detailed_changes) 기반 생성 + module_details 보완
    def _norm_txt(s: Any) -> str:
        return re.sub(r"\s+", " ", str(s or "").strip().upper())

    def _rule_based_module_from_change(dc: dict) -> str:
        scope = _norm_txt(f"{dc.get('part') or ''} {dc.get('change_detail') or ''} {dc.get('change_reason') or ''}")
        if not scope:
            return ""

        # 제품군 공통 규칙: 키워드 점수 높은 모듈을 우선 사용
        rules = {
            "Door Assembly": ["DOOR", "CAMERA", "HARNESS", "HINGE", "TRAY", "FRAME"],
            "Cavity": ["CAVITY", "CONV", "FAN", "MOTOR", "BLADE", "INLET", "OUTLET", "AIR"],
            "Controller Assembly": ["CONTROLLER", "LCD", "PANEL", "PCB", "UI", "OS", "S/W", "SW"],
        }
        best_mod = ""
        best_score = 0
        for mod, kws in rules.items():
            score = 0
            for kw in kws:
                if kw in scope:
                    score += 2
            # 강한 앵커 키워드 가중치
            if mod == "Door Assembly" and any(k in scope for k in ["HARNESS", "TRAY", "HINGE", "CAMERA"]):
                score += 2
            if mod == "Cavity" and any(k in scope for k in ["CONV", "FAN", "MOTOR", "BLADE"]):
                score += 2
            if mod == "Controller Assembly" and any(k in scope for k in ["LCD", "PCB", "PANEL", "UI", "OS"]):
                score += 2
            if score > best_score:
                best_mod = mod
                best_score = score
        return best_mod if best_score >= 3 else ""

    def _pick_module_from_summary(dc: dict) -> str:
        part = _norm_txt(dc.get("part"))
        detail = _norm_txt(dc.get("change_detail"))
        scope = f"{part} {detail}"
        stop_tokens = {"ASSY", "ASSEMBLY", "PART", "MODULE", "변경", "신규", "추가", "기구", "제어"}
        weak_tokens = {"COVER", "구조", "적용", "변경", "추가", "형상", "HOLE", "SIZE"}
        # 파트명 토큰을 더 신뢰하고, 일반 단어는 제외
        part_toks = [x for x in re.split(r"[^A-Z0-9가-힣]+", part) if len(x) >= 3 and x not in stop_tokens and x not in weak_tokens]
        detail_toks = [x for x in re.split(r"[^A-Z0-9가-힣]+", detail) if len(x) >= 4 and x not in stop_tokens and x not in weak_tokens]
        toks = part_toks + detail_toks
        best_mod = ""
        best_score = 0
        for it in module_summary_hints:
            mod = str(it.get("module") or "").strip()
            txt = _norm_txt(it.get("text"))
            if not mod or not txt:
                continue
            score = 0
            for tk in toks:
                if tk in txt:
                    score += 2
            if score > best_score:
                best_mod = mod
                best_score = score
        return best_mod if best_score >= 2 and len(toks) > 0 else ""

    def _pick_best_detail(dc: dict, preferred_module: str = "") -> dict | None:
        part = _norm_txt(dc.get("part"))
        detail = _norm_txt(dc.get("change_detail"))
        scope = f"{part} {detail}"
        best = None
        best_score = -1
        stop_tokens = {"ASSY", "ASSEMBLY", "PART", "MODULE", "변경", "신규", "추가", "기구", "제어"}

        # 부품명을 토큰화 (2자 이상)
        part_toks = [x for x in re.split(r"[^A-Z0-9가-힣]+", part) if len(x) >= 2 and x not in stop_tokens]

        for md in module_details:
            md_module = _norm_txt(md.get("module"))
            blob = " ".join([
                md_module,
                _norm_txt(md.get("sub_module")),
                _norm_txt((md.get("base") or {}).get("description")),
                _norm_txt((md.get("new") or {}).get("description")),
                _norm_txt(md.get("remark")),
            ])
            score = 0
            if preferred_module and _norm_txt(preferred_module) == md_module:
                score += 4
            # module명이 본문에 직접 등장하면 가중치 부여
            mod_toks = [x for x in re.split(r"[^A-Z0-9가-힣]+", md_module) if len(x) >= 3 and x not in stop_tokens]
            for mt in mod_toks:
                if mt in scope:
                    score += 2

            # 부품명 토큰: 3자 이상 +3, 2자 +1 (강화)
            for tok in part_toks:
                if tok in blob:
                    score += (3 if len(tok) >= 3 else 1)

            # 변경상세 토큰: 3자 이상만
            detail_toks = [x for x in re.split(r"[^A-Z0-9가-힣]+", detail) if len(x) >= 3 and x not in stop_tokens]
            for tok in detail_toks:
                if tok in blob:
                    score += 1

            if score > best_score:
                best = md
                best_score = score
        return best if best_score >= 1 else None

    def _infer_change_type(detail_text: Any, reason_text: Any, part_text: Any) -> str:
        blob = " ".join([str(detail_text or ""), str(reason_text or ""), str(part_text or "")]).upper()
        if any(k in blob for k in ["삭제", "제거", "미적용", "DELETED", "REMOVE"]):
            return "삭제"
        if any(k in blob for k in ["추가", "신규", "ADD", "NEW"]):
            return "NEW"
        return "Changing"

    change_points_for_bom = []
    for i, dc in enumerate(detailed_changes, 1):
        rule_mod = _rule_based_module_from_change(dc)
        summary_mod = _pick_module_from_summary(dc)
        preferred_mod = rule_mod or summary_mod
        md = _pick_best_detail(dc, preferred_module=preferred_mod)
        part = str(dc.get("part") or "정보 없음")
        detail = str(dc.get("change_detail") or "정보 없음")
        reason = str(dc.get("change_reason") or "정보 없음")
        concern = str(dc.get("concern") or "내용 없음")

        base_pno = "TBD"
        new_pno = "TBD"
        is_shared = False
        shared_src = ""
        related_parts: list[str] = []
        remark = ""
        discipline = str(dc.get("discipline") or dc.get("category") or "정보 없음").strip() or "정보 없음"
        module_name = "정보 없음"

        if preferred_mod:
            module_name = preferred_mod

        if md:
            md_module = str(md.get("module") or "").strip()
            if md_module:
                module_name = md_module
            base_pno = str((md.get("base") or {}).get("part_no") or "TBD")
            new_pno = str((md.get("new") or {}).get("part_no") or "TBD")
            is_shared = bool((md.get("new") or {}).get("is_shared"))
            shared_src = str((md.get("new") or {}).get("shared_with") or "")
            remark = str(md.get("remark") or "")
            rp = str(md.get("sub_module") or "").strip()
            if rp and _norm_txt(rp) != _norm_txt(part):
                related_parts.append(rp)

        # 구조 Assembly를 못 찾았을 때만 discipline(기구/제어)을 fallback으로 사용
        if module_name in ("", "정보 없음"):
            module_name = discipline

        desc = f"{part}: {detail}" if part and detail else (detail or part or "정보 없음")
        ctype = _infer_change_type(detail, reason, part)
        tags = list(dc.get("tags") or [])
        if is_shared and "공용" not in tags:
            tags.append("공용")

        row = {
            "id": i,
            "description": desc,
            "module": module_name,
            "discipline": discipline,
            "type": ctype,
            "base_part_no": base_pno,
            "new_part_no": new_pno,
            "is_shared_part": is_shared,
            "shared_source": shared_src,
            "related_parts": related_parts or ["정보 없음"],
            "concerns": [concern] if concern else ["내용 없음"],
            "tags": tags,
            "evidence": {
                "part": part,
                "change_reason": reason,
                "detail_src": dc.get("_src") or {},
                "discipline_src": "main_change_table.category",
                "module_detail_src": (md or {}).get("_src") if md else {},
            },
        }
        ui_type, ui_remark = _build_ui_meta_for_change_point(row)
        row["ui_type"] = ui_type
        row["ui_remark"] = ui_remark
        change_points_for_bom.append(row)

    # 최후 안전망: O/X 체크표 잔여 행 제거
    cp_filtered = []
    for r in change_points_for_bom:
        mod = str(r.get("module") or "")
        desc = re.sub(r"\s+", "", str(r.get("description") or "").upper())
        if "과전압" in mod:
            discarded.append({
                "why": "cp_drop_blacklist_module",
                "slide": None,
                "table_id": "final_cp",
                "sample_text": f"{mod} | {r.get('description','')}",
            })
            continue
        if desc in {"O", "X", "(O,X)", "OX"}:
            discarded.append({
                "why": "cp_drop_checkbox_token",
                "slide": None,
                "table_id": "final_cp",
                "sample_text": f"{mod} | {r.get('description','')}",
            })
            continue
        cp_filtered.append(r)

    return {
        "project_meta": project_meta,
        "module_summary": module_summary,
        "detailed_changes": detailed_changes,
        "module_details": module_details,
        "risk_review": risk_review,
        "change_points_for_bom": cp_filtered,
        "attachment_presence": image_presence,
        "slide_roles": slide_role_details,
        "discarded": discarded,
    }
