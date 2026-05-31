from __future__ import annotations

import re
from pathlib import Path

RELEVANT_KEYWORDS = ["변경점", "변경내역", "변경 내역", "유첨1", "유첨3", "개발 변경", "PART", "BASE", "NEW"]


def _ppt_text_from_shape(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    chunks = []
    for p in shape.text_frame.paragraphs:
        run_text = []
        for r in p.runs:
            try:
                if getattr(r.font, "strike", False):
                    continue
            except Exception:
                pass
            run_text.append(r.text or "")
        line = "".join(run_text).strip()
        if line:
            chunks.append(line)
    return "\n".join(chunks).strip()


def _table_to_markdown(table) -> str:
    rows = []
    for row in table.rows:
        cells = [re.sub(r"\s+", " ", cell.text or "").strip() for cell in row.cells]
        non_empty = [c for c in cells if c]
        if non_empty:
            rows.append("| " + " | ".join(cells) + " |")
    return "\n".join(rows)


def _is_relevant(title: str, content: str) -> bool:
    blob = f"{title}\n{content}".upper()
    return any(kw.upper() in blob for kw in RELEVANT_KEYWORDS)


def extract_slides(pptx_path: str | Path) -> list[dict]:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise RuntimeError("python-pptx가 필요합니다: pip install python-pptx") from e

    pptx_path = Path(pptx_path)
    prs = Presentation(pptx_path)
    results = []

    for i, slide in enumerate(prs.slides, 1):
        texts = []
        tables = []

        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                md = _table_to_markdown(shape.table)
                if md:
                    tables.append(md)
            elif getattr(shape, "has_text_frame", False):
                txt = _ppt_text_from_shape(shape)
                if txt:
                    texts.append(txt)

        title = texts[0] if texts else ""
        content_parts = texts + tables
        content = "\n\n".join(content_parts)

        if not _is_relevant(title, content):
            continue

        results.append({
            "slide": i,
            "title": title,
            "content": content,
            "has_table": bool(tables),
            "source_pptx": pptx_path.name,
        })

    return results


def load_all_pptx(input_dir: str | Path) -> list[dict]:
    input_dir = Path(input_dir)
    slides = []
    for pptx_path in sorted(input_dir.glob("*.pptx")):
        if pptx_path.name.startswith("~$"):
            continue
        try:
            slides.extend(extract_slides(pptx_path))
        except Exception as e:
            print(f"[WARNING] {pptx_path.name} 로드 실패: {e}")
    return slides
